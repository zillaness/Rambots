"""
file: build_print_tracker_v1.0.py
version: 1.0
author: Samuel Cao
created: 2026-10-08
last_updated: 2026-10-08
description: Builds the Print Tracker workbook (.xlsx); uploaded to Google Drive it converts to a native Google Sheet.
ai_update: Update last_updated and version. Rename file to match. Append changelog at bottom.

Usage:
    python3 tools/print_tracker/build_print_tracker_v1.0.py --out print_tracker_v1.0.xlsx
    python3 tools/print_tracker/build_print_tracker_v1.0.py --blank --out print_tracker_blank_v1.0.xlsx

Several formulas (QUERY, MINIFS, MAXIFS, INDIRECT) are written for Google Sheets. Excel or LibreOffice
will show errors in those cells; Google Sheets evaluates them after import.
"""

import argparse
import re
import zipfile
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

VERSION = "1.0"
AUTHOR = "Samuel Cao"
DATE = "2026-10-08"

# Rows pre-filled with formulas. Kept small so the file uploads cheaply; the Apps Script copies the
# formulas to any row past these when someone types a Part there or Fill gaps writes there.
PRINTS_FORMULA_ROWS = 120
PARTS_FORMULA_ROWS = 40
FORMAT_ROWS = 2000  # rows covered by dropdowns and conditional formatting
DEFAULT_BATCH = 10
STALE_DAYS = 3

# Rambot Red header with white text; pale fills mark formula and fill-me-in cells.
RED = "A31F36"
HEADER_FILL = PatternFill("solid", fgColor=RED)
HEADER_FONT = Font(bold=True, color="FFFFFF")
FORMULA_FILL = PatternFill("solid", fgColor="F3F3F3")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
SECTION_FONT = Font(bold=True, size=12, color=RED)

HIDE_ZERO_INT = "0;-0;;@"
HIDE_ZERO_DEC = "0.0;-0.0;;@"

LISTS = {
    "Status": ["Unclaimed", "Claimed", "Printing", "Done", "Delivered", "Failed"],
    "Priority": ["Must", "Stretch"],
    "Material": ["PLA", "PLA+", "PLA Silk", "PETG", "ABS", "ASA", "TPU", "PLA-CF", "PETG-CF", "PA-CF", "Other"],
    "Color": [
        "Assorted", "Black", "White", "Grey", "Red", "Orange", "Yellow", "Green", "Blue",
        "Purple", "Pink", "Brown", "Translucent", "Glow", "Silk", "Other",
    ],
    "Printer model": ["A1 mini", "A1", "P1P", "P1S", "X1C", "X1E", "H2D", "Other"],
    "Owner": ["Team", "Personal"],
    "Yes": ["Yes"],
    "Projects": [],
}

CARNIVAL = "Ramathon Carnival"
MULTI_PART_NOTE = "Confirm the design: if it prints as several pieces, give each piece its own Parts row under the same Kit."


@dataclass
class Part:
    kit: str
    part: str
    priority: str
    needed: int
    notes: str = ""
    project: str = CARNIVAL
    per_kit: int = 1
    per_plate: int = DEFAULT_BATCH  # placeholder until someone slices a plate
    material: str = "PLA"
    color: str = "Assorted"


LAUNCHER = "Pull copter launcher"
SEED_PARTS = [
    Part(LAUNCHER, "Launcher handle", "Must", 80),
    Part(LAUNCHER, "Launcher gear", "Must", 80),
    Part(LAUNCHER, "Launcher shaft", "Must", 80),
    Part(LAUNCHER, "Launcher bearing", "Must", 80),
    Part(LAUNCHER, "Launcher sleeve", "Must", 80),
    Part("Ripcord", "Ripcord", "Must", 120),
    Part("Basic propeller", "Basic propeller", "Must", 120),
    Part("Fancy propeller", "Fancy propeller", "Must", 60),
    Part("Flexi rex", "Flexi rex", "Stretch", 30),
    Part("Infinity cube", "Infinity cube", "Stretch", 30, MULTI_PART_NOTE),
    Part("Fidget clicker", "Fidget clicker", "Stretch", 30, MULTI_PART_NOTE),
]


@dataclass
class Col:
    header: str
    width: int
    formula: str = ""  # template; {r} is the row number
    number_format: str = ""
    note: str = ""
    hidden: bool = False


def lookup(col):
    """Text from the matching Parts row, via the hidden _row helper in Prints column R."""
    return f'=IF($R{{r}}="","",INDEX(Parts!${col}:${col},$R{{r}})&"")'


def per_copy(col):
    """This row's share of a per-plate figure on Parts (hours or grams). 0 when unknown."""
    return (
        f'=IF($R{{r}}="",0,IFERROR(IF(INDEX(Parts!${col}:${col},$R{{r}})="",0,'
        f"$B{{r}}/INDEX(Parts!$G:$G,$R{{r}})*INDEX(Parts!${col}:${col},$R{{r}})),0))"
    )


def needs_flag(col, label):
    return f'IF(INDEX(Parts!${col}:${col},$R{{r}})="Yes","{label} ","")'


PRINTS_COLS = [
    Col("Part", 22, note="Pick from the list. Parts are defined on the Parts tab."),
    Col("Qty", 6, note="Copies in this batch. Change it if your plate fits more or fewer."),
    Col("Priority", 9, lookup("D")),
    Col("Material", 10, lookup("H")),
    Col(
        "Needs",
        20,
        '=IF($R{r}="","",TRIM('
        + "&".join([needs_flag("J", "Enclosure"), needs_flag("K", "Hardened"), needs_flag("L", "Supports"), needs_flag("M", ">180mm")])
        + "))",
        note="Printer requirements, from the Parts tab.",
    ),
    Col("Color", 11, note="Pick the color you're printing when you claim the row."),
    Col("Claimed by", 14, note="Pick your name to claim. Clear it to give the row back."),
    Col("Status", 11, note="Unclaimed, Claimed, Printing, Done (printed), Delivered (handed in), Failed (not retrying)."),
    Col("Est. hours", 9, per_copy("N"), HIDE_ZERO_DEC),
    Col("Est. grams", 9, per_copy("O"), HIDE_ZERO_INT),
    Col(
        "File",
        8,
        '=IF($R{r}="","",IF(INDEX(Parts!$P:$P,$R{r})="","",HYPERLINK(INDEX(Parts!$P:$P,$R{r}),"Open")))',
    ),
    Col("Slicing notes", 30, lookup("Q")),
    Col("Claimed on", 11, number_format="mmm d", note="Filled in by the script when a name is picked."),
    Col("Status changed", 13, number_format="mmm d h:mm", note="Filled in by the script on every status change. Rows go orange when this is too old."),
    Col("Notes", 24),
    Col("Project", 16, lookup("A")),
    Col("Kit", 18, lookup("B")),
    Col("_row", 5, '=IF($A{r}="","",IFERROR(MATCH($A{r},Parts!$C:$C,0),""))', hidden=True),
]

KIT = "$A:$A,$A{r},$B:$B,$B{r}"
PARTS_COLS = [
    Col("Project", 18, note="Add new project names to Lists > Projects first."),
    Col("Kit", 20, note="Parts that get assembled together share a Kit name. A standalone item is its own kit."),
    Col("Part", 22, note="Must be unique. This is the name people pick on Prints."),
    Col("Priority", 9, note="Must rows are planned and listed first."),
    Col("Qty needed", 9),
    Col("Per kit", 7, note="Copies of this part in one finished kit."),
    Col("Per plate", 8, note="Copies on one plate. Slice one plate in Bambu Studio, enter it here, then run Print Tracker > Re-split unclaimed rows."),
    Col("Material", 9),
    Col("Color", 11),
    Col("Enclosure", 9),
    Col("Hardened nozzle", 9),
    Col("Supports", 9),
    Col("Needs >180 mm bed", 9, note="Yes rules out the A1 mini."),
    Col("Hours per plate", 9, note="From Bambu Studio after slicing one full plate."),
    Col("Grams per plate", 9, note="From Bambu Studio after slicing one full plate."),
    Col("File link", 14, note="Bambu Studio project 3MF (unsliced) in the team Drive folder."),
    Col("Slicing notes", 30),
    Col("In rows", 8, '=IF($C{r}="","",SUMIFS(Prints!$B:$B,Prints!$A:$A,$C{r},Prints!$H:$H,"<>Failed"))', note="Copies covered by Prints rows (Failed rows don't count)."),
    Col("Gap", 8, '=IF($C{r}="","",$E{r}-$R{r})', note="Positive: copies with no row yet, run Fill gaps. Negative: rows cover more than needed."),
    Col("Done", 8, '=IF($C{r}="","",SUMIFS(Prints!$B:$B,Prints!$A:$A,$C{r},Prints!$H:$H,"Done")+SUMIFS(Prints!$B:$B,Prints!$A:$A,$C{r},Prints!$H:$H,"Delivered"))', note="Printed copies (Done + Delivered)."),
    Col("Delivered", 8, '=IF($C{r}="","",SUMIFS(Prints!$B:$B,Prints!$A:$A,$C{r},Prints!$H:$H,"Delivered"))'),
    Col("% done", 8, '=IF(N($E{r})>0,MIN(1,$T{r}/$E{r}),"")', "0%"),
    Col("Kits needed", 8, '=IF(N($F{r})>0,$E{r}/$F{r},0)', HIDE_ZERO_INT),
    Col("Kits done", 8, '=IF(N($F{r})>0,INT($T{r}/$F{r}),0)', HIDE_ZERO_INT, note="Finished kits this part allows."),
    Col(
        "Bottleneck",
        11,
        f'=IF(OR($C{{r}}="",COUNTIFS({KIT})<2),"",IF(MAXIFS($X:$X,{KIT})=MINIFS($X:$X,{KIT}),"",'
        f'IF($X{{r}}=MINIFS($X:$X,{KIT}),"Bottleneck","")))',
        note="The part holding its kit back.",
    ),
    Col("Hours left", 9, '=IF(OR($C{r}="",N($G{r})=0,N($N{r})=0),"",ROUNDUP(MAX(0,$E{r}-$T{r})/$G{r},0)*$N{r})', "0.0"),
    Col("Grams left", 9, '=IF(OR($C{r}="",N($G{r})=0,N($O{r})=0),"",ROUNDUP(MAX(0,$E{r}-$T{r})/$G{r},0)*$O{r})', "0"),
]

PRINTERS_COLS = [
    Col("Printer", 18, note="A name people will recognize, e.g. 'Team P1S #2' or 'Ava's A1'."),
    Col("Model", 10),
    Col("Owner", 9),
    Col("Currently with", 14, note="Who has it now. For a personal printer, the owner."),
    Col("Since", 11, number_format="mmm d", note="Filled in by the script when Currently with changes."),
    Col("Enclosed", 9),
    Col("Hardened nozzle", 9, note="Whatever is installed right now, not the stock nozzle."),
    Col("Max print size (mm)", 10),
    Col("Notes", 30),
]

PEOPLE_COLS = [Col("Name", 18), Col("Contact", 22), Col("Notes", 30)]


def style_header(ws, cols):
    for i, col in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i, value=col.header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        if col.note:
            cell.comment = Comment(col.note, AUTHOR, width=260, height=90)
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = col.width
        if col.hidden:
            ws.column_dimensions[letter].hidden = True
    ws.row_dimensions[1].height = 32


def fill_formulas(ws, cols, last_row):
    for i, col in enumerate(cols, start=1):
        if not (col.formula or col.number_format):
            continue
        for r in range(2, last_row + 1):
            cell = ws.cell(row=r, column=i)
            if col.formula:
                cell.value = col.formula.format(r=r)
                cell.fill = FORMULA_FILL
            if col.number_format:
                cell.number_format = col.number_format


def col_letter(cols, header):
    return get_column_letter([c.header for c in cols].index(header) + 1)


def add_list(ws, cols, header, source, last_row=FORMAT_ROWS):
    letter = col_letter(cols, header)
    dv = DataValidation(type="list", formula1=source, allow_blank=True, showErrorMessage=True, errorStyle="warning")
    dv.add(f"{letter}2:{letter}{last_row}")
    ws.add_data_validation(dv)


def lists_range(name):
    letter = get_column_letter(list(LISTS).index(name) + 1)
    return f"Lists!${letter}$2:${letter}$60"


def plan_rows(parts):
    """Same rule as planGapRows in the Apps Script: Must first, split into Per plate batches."""
    rows = []
    for p in sorted(parts, key=lambda p: p.priority != "Must"):
        left = p.needed
        while left > 0:
            qty = min(p.per_plate, left)
            rows.append((p.part, qty))
            left -= qty
    return rows


def build_prints(wb, parts):
    ws = wb.create_sheet("Prints")
    cols = PRINTS_COLS
    style_header(ws, cols)
    fill_formulas(ws, cols, PRINTS_FORMULA_ROWS + 1)
    for r, (part, qty) in enumerate(plan_rows(parts), start=2):
        ws.cell(row=r, column=1, value=part)
        ws.cell(row=r, column=2, value=qty)
        ws.cell(row=r, column=8, value="Unclaimed")
    ws.freeze_panes = "B2"

    add_list(ws, cols, "Part", "Parts!$C$2:$C$500")
    add_list(ws, cols, "Color", lists_range("Color"))
    add_list(ws, cols, "Claimed by", "People!$A$2:$A$300")
    add_list(ws, cols, "Status", lists_range("Status"))

    # First matching rule wins in Google Sheets, so stale goes first.
    span = f"A2:Q{FORMAT_ROWS}"
    rules = [
        ('AND(OR($H2="Claimed",$H2="Printing"),$N2<>"",TODAY()-$N2>INDIRECT("Summary!$B$4"))', "F9CB9C"),
        ('$H2="Failed"', "F4CCCC"),
        ('$H2="Delivered"', "B6D7A8"),
        ('$H2="Done"', "D9EAD3"),
        ('$H2="Printing"', "FFF2CC"),
        ('$H2="Claimed"', "CFE2F3"),
    ]
    for formula, color in rules:
        ws.conditional_formatting.add(span, FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=color)))


def build_parts(wb, parts):
    ws = wb.create_sheet("Parts")
    cols = PARTS_COLS
    style_header(ws, cols)
    fill_formulas(ws, cols, PARTS_FORMULA_ROWS + 1)
    for r, p in enumerate(parts, start=2):
        values = [p.project, p.kit, p.part, p.priority, p.needed, p.per_kit, p.per_plate, p.material, p.color]
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=v)
        ws.cell(row=r, column=17, value=p.notes or None)
        # Placeholder plate size and the numbers only slicing can give: flag them yellow.
        for header in ("Per plate", "Hours per plate", "Grams per plate", "File link"):
            ws[f"{col_letter(cols, header)}{r}"].fill = INPUT_FILL
    ws.freeze_panes = "D2"

    last = PARTS_FORMULA_ROWS + 100
    add_list(ws, cols, "Project", lists_range("Projects"), last)
    add_list(ws, cols, "Priority", lists_range("Priority"), last)
    add_list(ws, cols, "Material", lists_range("Material"), last)
    add_list(ws, cols, "Color", lists_range("Color"), last)
    for header in ("Enclosure", "Hardened nozzle", "Supports", "Needs >180 mm bed"):
        add_list(ws, cols, header, lists_range("Yes"), last)

    gap = col_letter(cols, "Gap")
    ws.conditional_formatting.add(f"{gap}2:{gap}{last}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="F9CB9C")))
    ws.conditional_formatting.add(f"{gap}2:{gap}{last}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="CC0000", bold=True)))
    neck = col_letter(cols, "Bottleneck")
    ws.conditional_formatting.add(f"{neck}2:{neck}{last}", CellIsRule(operator="equal", formula=['"Bottleneck"'], font=Font(color="CC0000", bold=True)))


def build_summary(wb, project):
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 34
    for letter in "BCDEFGHIJ":
        ws.column_dimensions[letter].width = 14

    ws["A1"] = "Print tracker summary"
    ws["A1"].font = Font(bold=True, size=14, color=RED)

    settings = [
        ("Project (blank = all projects)", project or None, None),
        ("Deadline", None, "yyyy-mm-dd"),
        ("Stale after (days)", STALE_DAYS, None),
        ("Default batch size", DEFAULT_BATCH, None),
    ]
    for r, (label, value, fmt) in enumerate(settings, start=2):
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=r, column=2, value=value)
        cell.fill = INPUT_FILL
        if fmt:
            cell.number_format = fmt
    ws["B3"].comment = Comment("Type the event date. Days left and the capacity line use it.", AUTHOR)
    ws["B5"].comment = Comment("Batch size Fill gaps uses for parts with no Per plate yet.", AUTHOR)
    dv = DataValidation(type="list", formula1=lists_range("Projects"), allow_blank=True)
    dv.add("B2")
    ws.add_data_validation(dv)
    for name, ref in (("Deadline", "$B$3"), ("StaleDays", "$B$4"), ("DefaultBatch", "$B$5")):
        wb.defined_names[name] = DefinedName(name, attr_text=f"Summary!{ref}")

    proj = 'IF($B$2="","*",$B$2)'
    by = lambda col: f"Parts!${col}:${col},Parts!$D:$D,{{c}}$7,Parts!$A:$A,{proj}"  # noqa: E731
    for c in ("B", "C"):
        ws[f"{c}7"] = "Must" if c == "B" else "Stretch"
        ws[f"{c}7"].fill = HEADER_FILL
        ws[f"{c}7"].font = HEADER_FONT
    ws["A7"].fill = HEADER_FILL
    rows = [
        ("Copies needed", "=SUMIFS(" + by("E") + ")", "0"),
        ("Copies printed", "=SUMIFS(" + by("T") + ")", "0"),
        ("% printed", "=IF({c}8>0,{c}9/{c}8,\"\")", "0%"),
        ("Copies with no row yet", "=SUMIFS(" + by("S") + ',Parts!$S:$S,">0")', "0"),
        ("Unclaimed rows", f'=COUNTIFS(Prints!$C:$C,{{c}}$7,Prints!$H:$H,"Unclaimed",Prints!$P:$P,{proj})', "0"),
        ("Print hours left", "=SUMIFS(" + by("Z") + ")", "0.0"),
        ("Parts missing hours per plate", f'=COUNTIFS(Parts!$D:$D,{{c}}$7,Parts!$A:$A,{proj},Parts!$N:$N,"")', "0"),
        ("Filament left (g)", "=SUMIFS(" + by("AA") + ")", "#,##0"),
        ("Parts missing grams per plate", f'=COUNTIFS(Parts!$D:$D,{{c}}$7,Parts!$A:$A,{proj},Parts!$O:$O,"")', "0"),
    ]
    for r, (label, formula, fmt) in enumerate(rows, start=8):
        ws.cell(row=r, column=1, value=label)
        for c in ("B", "C"):
            ws[f"{c}{r}"] = formula.format(c=c)
            ws[f"{c}{r}"].number_format = fmt

    ws["A18"] = "Capacity (Must only)"
    ws["A18"].font = SECTION_FONT
    capacity = [
        ("Printers listed", "=COUNTA(Printers!$A$2:$A$500)", "0"),
        ("Days left", '=IF($B$3="","",MAX(0,$B$3-TODAY()))', "0"),
        (
            "Print hours per printer per day needed",
            '=IF($B$3="","Set the deadline",IF(B19=0,"Add printers on the Printers tab",'
            'IF(B14>0,"Fill hours per plate for "&B14&" part(s)",IF(B20=0,"Deadline reached",B13/B19/B20))))',
            "0.0",
        ),
        (
            "Stale claims (all projects)",
            '=IFERROR(SUMPRODUCT(((Prints!$H$2:$H$2000="Claimed")+(Prints!$H$2:$H$2000="Printing"))'
            '*(Prints!$N$2:$N$2000<>"")*(TODAY()-Prints!$N$2:$N$2000>$B$4)),0)',
            "0",
        ),
        ("Failed rows (all projects)", '=COUNTIF(Prints!$H:$H,"Failed")', "0"),
    ]
    for r, (label, formula, fmt) in enumerate(capacity, start=19):
        ws.cell(row=r, column=1, value=label)
        ws[f"B{r}"] = formula
        ws[f"B{r}"].number_format = fmt

    parts_filter = '&IF($B$2="",""," and A = \'"&$B$2&"\'")&'
    prints_filter = '&IF($B$2="",""," and P = \'"&$B$2&"\'")&'
    tables = [
        (
            25,
            "Kits (finished kits = copies of the bottleneck part)",
            '=IFERROR(QUERY(Parts!$A$1:$AA$500,"select A, D, B, max(W), min(X) where C <> \'\'"'
            + parts_filter
            + '" group by A, D, B label A \'Project\', D \'Priority\', B \'Kit\', max(W) \'Kits needed\', min(X) \'Kits done\'",1),"No parts yet")',
        ),
        (
            52,
            "Who has what (copies by status)",
            '=IFERROR(QUERY(Prints!$A$1:$R$2000,"select G, sum(B) where G <> \'\'"'
            + prints_filter
            + '" group by G pivot H",1),"Nobody has claimed anything yet")',
        ),
        (
            82,
            "Filament to hand out (rows Claimed or Printing)",
            '=IFERROR(QUERY(Prints!$A$1:$R$2000,"select G, D, F, sum(J) where G <> \'\' and (H = \'Claimed\' or H = \'Printing\')"'
            + prints_filter
            + '" group by G, D, F label G \'Person\', D \'Material\', F \'Color\', sum(J) \'Grams\'",1),"Nothing claimed yet")',
        ),
        (
            122,
            "Color mix (printed copies)",
            '=IFERROR(QUERY(Prints!$A$1:$R$2000,"select A, sum(B) where (H = \'Done\' or H = \'Delivered\')"'
            + prints_filter
            + '" group by A pivot F",1),"Nothing printed yet")',
        ),
    ]
    for row, title, formula in tables:
        ws.cell(row=row - 1, column=1, value=title).font = SECTION_FONT
        ws.cell(row=row, column=1, value=formula)


def build_simple(wb, title, cols, lists):
    ws = wb.create_sheet(title)
    style_header(ws, cols)
    fill_formulas(ws, cols, 100)
    for header, source in lists:
        add_list(ws, cols, header, source, 300)
    ws.freeze_panes = "A2"
    return ws


HOW_TO = [
    ("Print Tracker", "title"),
    ("Claiming prints", "section"),
    ("1. On Prints, find an Unclaimed row your printer can do. Check Needs and Material. Must rows come first.", ""),
    ("2. Pick your name in Claimed by. The date fills in and Status turns Claimed.", ""),
    ("3. Pick the Color you'll print. Ask whoever hands out filament for the spool you need.", ""),
    ("4. Update Status as you go: Printing, Done (printed), Delivered (handed in at the shop).", ""),
    ("5. Plate fits more or fewer copies? Change Qty. A lead runs Fill gaps to cover the rest.", ""),
    ("6. Print failed and you're not retrying? Set Failed. Fill gaps puts those copies back up for grabs.", ""),
    ("7. Can't finish it? Clear your name and the row goes back to Unclaimed.", ""),
    ("Rows turn orange when a Claimed or Printing row hasn't changed in 3 days (set on Summary).", ""),
    ("To see only your rows, use Data > Create a filter view. It doesn't change what anyone else sees.", ""),
    ("Running a project (leads)", "section"),
    ("New project: add its name to Lists > Projects, then add its parts on Parts.", ""),
    ("Each part: Project, Kit, Part, Priority, Qty needed, Per kit, Per plate, Material, Color, file link. Then Print Tracker > Fill gaps.", ""),
    ("After slicing one plate in Bambu Studio: fill Per plate, Hours per plate and Grams per plate, then Print Tracker > Re-split unclaimed rows.", ""),
    ("Link files as Bambu Studio project 3MFs (oriented, settings saved, not sliced) from the team Drive folder.", ""),
    ("Add people on People and printers on Printers so the dropdowns and capacity line work.", ""),
    ("Grey columns are formulas. Don't type in them. Add Prints rows in the first empty row; don't insert rows mid-table.", ""),
    ("Set the Deadline on Summary. 'Print hours per printer per day' tells you early whether the plan fits.", ""),
    ("One-time setup", "section"),
    ("Extensions > Apps Script. Replace the contents of Code.gs with print_tracker_v1.0.gs from the Rambots repo (tools/print_tracker/). Save.", ""),
    ("Reload the sheet. A Print Tracker menu appears. The first menu click asks for permission to edit this spreadsheet only.", ""),
    ("Test once from the Sheets phone app: claim a row and check that Claimed on fills in.", ""),
    ("About", "section"),
    (f"Version {VERSION}. Created {DATE} by {AUTHOR}. Built by tools/print_tracker/build_print_tracker_v{VERSION}.py.", ""),
    (f"Changelog: v1.0 ({DATE}) Initial release.", ""),
]


def build_how_to(wb):
    ws = wb.create_sheet("How to use")
    ws.column_dimensions["A"].width = 110
    for r, (text, kind) in enumerate(HOW_TO, start=1):
        cell = ws.cell(row=r, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "title":
            cell.font = Font(bold=True, size=14, color=RED)
        elif kind == "section":
            cell.font = SECTION_FONT


def build_lists(wb, project):
    ws = wb.create_sheet("Lists")
    LISTS["Projects"] = [project] if project else []
    for c, (name, values) in enumerate(LISTS.items(), start=1):
        head = ws.cell(row=1, column=c, value=name)
        head.fill = HEADER_FILL
        head.font = HEADER_FONT
        ws.column_dimensions[get_column_letter(c)].width = 14
        for r, v in enumerate(values, start=2):
            ws.cell(row=r, column=c, value=v)


def build(blank):
    parts = [] if blank else SEED_PARTS
    project = None if blank else CARNIVAL
    wb = Workbook()
    wb.remove(wb.active)
    build_prints(wb, parts)
    build_summary(wb, project)
    build_parts(wb, parts)
    build_simple(
        wb,
        "Printers",
        PRINTERS_COLS,
        [
            ("Model", lists_range("Printer model")),
            ("Owner", lists_range("Owner")),
            ("Currently with", "People!$A$2:$A$300"),
            ("Enclosed", lists_range("Yes")),
            ("Hardened nozzle", lists_range("Yes")),
        ],
    )
    build_simple(wb, "People", PEOPLE_COLS, [])
    build_how_to(wb)
    build_lists(wb, project)
    return wb


FORMULA_CELL = re.compile(r'<c r="([A-Z]+)(\d+)"((?: s="\d+")?)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>')


def share_formulas(path, sheet_names):
    """
    Rewrites each column's per-row formulas as one Excel shared formula. Every row's formula comes
    from the same template, so this is lossless, and it cuts the file size several times over.
    """
    with zipfile.ZipFile(path) as z:
        entries = {name: z.read(name) for name in z.namelist()}

    for sheet in sheet_names:
        name = f"xl/worksheets/{sheet}.xml"
        xml = entries[name].decode("utf-8")
        cells = FORMULA_CELL.findall(xml)
        first, last = {}, {}
        for col, row, _, _ in cells:
            first.setdefault(col, int(row))
            last[col] = int(row)
        ids = {col: i for i, col in enumerate(first)}

        def replace(m):
            col, row, style, formula = m.group(1), int(m.group(2)), m.group(3), m.group(4)
            si = ids[col]
            if row == first[col]:
                f = f'<f t="shared" ref="{col}{row}:{col}{last[col]}" si="{si}">{formula}</f>'
            else:
                f = f'<f t="shared" si="{si}"/>'
            return f'<c r="{col}{row}"{style}>{f}</c>'

        entries[name] = FORMULA_CELL.sub(replace, xml).encode("utf-8")

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for name, data in entries.items():
            z.writestr(name, data)


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--out", default=f"print_tracker_v{VERSION}.xlsx")
    parser.add_argument("--blank", action="store_true", help="no seed project, for starting a new tracker")
    args = parser.parse_args()
    wb = build(args.blank)
    wb.save(args.out)
    # openpyxl numbers worksheet files in tab order: sheet1.xml, sheet2.xml, ...
    share_formulas(args.out, [f"sheet{wb.sheetnames.index(t) + 1}" for t in ("Prints", "Parts")])
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()

# CHANGELOG
# v1.0 (2026-10-08): Initial release. Prints, Summary, Parts, Printers, People, How to use and Lists tabs,
#   seeded with the Ramathon Carnival parts.
