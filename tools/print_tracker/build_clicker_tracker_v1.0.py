"""
file: build_clicker_tracker_v1.0.py
version: 1.0
author: Samuel Cao
created: 2026-10-20
last_updated: 2026-10-20
description: Builds the clicker print tracker workbook (.xlsx): orders, printed components, print log, progress and a plate planner.
ai_update: Update last_updated and version. Rename file to match. Append changelog at bottom.

Usage:
    python3 tools/print_tracker/build_clicker_tracker_v1.0.py --out clicker_print_tracker_v1.0.xlsx

Uses only formulas that Excel (2019+), Google Sheets and LibreOffice all evaluate: SUMIF(S), COUNTIF(S),
INDEX/MATCH and MINIFS (written as _xlfn.MINIFS, as Excel stores it). No dynamic arrays, so the planner lists use a rank column instead of FILTER.
"""

import argparse
import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AUTHOR = "Samuel Cao"
VERSION = "1.0"
DATE = "2026-10-20"

COMPONENT_ROWS = 150  # rows pre-filled with formulas on Components
PROGRESS_ROWS = 60
LOG_ROWS = 1000
PLANNER_ROWS = 40
LOADED_FIRST, LOADED_LAST = 5, 12  # Plate planner A5:A12, up to 8 loaded colors

RED = "A31F36"
HEADER_FILL = PatternFill("solid", fgColor=RED)
HEADER_FONT = Font(bold=True, color="FFFFFF")
FORMULA_FILL = PatternFill("solid", fgColor="EEF2F8")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
TITLE_FONT = Font(bold=True, size=14, color=RED)
NOTE_FONT = Font(italic=True, color="5F6368")

# Order list, copied from the "Robotics internal clicker prints" Google Sheet (2026-10-20).
ORDERS = [
    ("dumpling/ xiao long bao", 12, 21),
    ("avocado", 4, 0),
    ("creeper", 5, 0),
    ("salmon nigiri", 8, 21),
    ("infinity cube", 4, 0),
    ("chocolate", 4, 0),
    ("dan tat", 15, 21),
    ("smores", 2, 0),
    ("oreo", 0, 0),
    ("hot dog bun", 4, 0),
    ("ramen", 0, 21),
    ("sushi", 0, 21),
]

CONFIRM = "Confirm: one multicolor print, or separate pieces?"
# (part, item, per item, colors, notes). Colors in one row print together as one multicolor print.
COMPONENTS = [
    ("Dumpling top", "dumpling/ xiao long bao", 1, ["white", "black", "red"], ""),
    ("Dumpling base", "dumpling/ xiao long bao", 1, ["tan"], ""),
    ("Avocado", "avocado", 1, ["green", "yellow green", "brown", "black"], CONFIRM),
    ("Creeper", "creeper", 1, ["green", "black"], CONFIRM),
    ("Salmon nigiri fish", "salmon nigiri", 1, ["orange"], "Split from rice because of '30 rice parts?'. Confirm"),
    ("Salmon nigiri rice", "salmon nigiri", 1, ["white"], "Same piece as Sushi rice? If so, merge the rows"),
    ("Infinity cube", "infinity cube", 1, ["green", "black"], "Confirm pieces per cube"),
    ("Chocolate", "chocolate", 1, ["brown"], ""),
    ("Dan tat", "dan tat", 1, ["yellow", "tan"], CONFIRM),
    ("Smores", "smores", 1, ["brown", "tan", "white"], CONFIRM),
    ("Oreo", "oreo", 1, ["black", "white"], CONFIRM),
    ("Hot dog", "hot dog bun", 1, ["red", "tan"], CONFIRM),
    ("Ramen bowl", "ramen", 1, ["red", "black"], ""),
    ("Ramen egg", "ramen", 1, ["yellow", "white"], "Yolk orange or yellow?"),
    ("Ramen soup", "ramen", 1, ["green", "pink", "tan", "brown"], "Scallions, meat, noodles, broth"),
    ("Ramen chopsticks", "ramen", 1, ["brown", "green"], "Or red instead of green"),
    ("Sushi nori", "sushi", 1, ["black"], "Tuna sushi, piece 1 of 3"),
    ("Sushi rice", "sushi", 1, ["white"], "Tuna sushi, piece 2 of 3"),
    ("Sushi fish", "sushi", 1, ["TBD"], "Tuna sushi, piece 3 of 3. Pick a color"),
]

PRINT_LOG = [
    (datetime.date(2026, 10, 20), "Sam", "Dumpling top", 18, 8, "", "From Sheet1: 18 tops, 8 B-stock"),
    (datetime.date(2026, 10, 20), "Sam", "Salmon nigiri rice", 30, 0, "", "From Sheet1: '30 rice parts?' Confirm"),
]

PRINTERS = [
    ("X1C + AMS", "Sam", 4, ""),
    ("P1S + AMS", "Sam", 4, ""),
    ("H2D + 2x AMS 2 Pro", "Sam", 8, ""),
    ("Snapmaker U1", "Sam", 4, "Toolchanger"),
    ("A1 mini", "Makerspace", 1, ""),
    ("A1 mini", "Makerspace", 1, ""),
    ("Elegoo Carbon Combo 2", "Makerspace", None, "How many colors?"),
    ("Elegoo Carbon Combo 2", "Makerspace", None, "How many colors?"),
    ("A1 mini + AMS lite", "Asher", 4, ""),
    ("A1", "Quinn", 1, ""),
    ("A1", "Ayden", 1, ""),
    ("Anycubic Kobra (large bed)", "Ryan", 1, "Model unconfirmed"),
]

COLORS = [
    "white", "black", "red", "orange", "yellow", "yellow green", "green", "blue", "purple",
    "pink", "tan", "brown", "grey", "TBD",
]


def header(ws, labels, widths, formula_from=None):
    for i, (label, width) in enumerate(zip(labels, widths), start=1):
        cell = ws.cell(row=1, column=i, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B2"


def dropdown(ws, cells, source):
    dv = DataValidation(type="list", formula1=source, allow_blank=True, showErrorMessage=True, errorStyle="warning")
    dv.add(cells)
    ws.add_data_validation(dv)


def build_orders(wb):
    ws = wb.create_sheet("Orders")
    header(ws, ["Item", "Internal", "Customer", "Total (= switches)"], [26, 10, 10, 12])
    for r in range(2, PROGRESS_ROWS + 2):
        ws[f"D{r}"] = f'=IF(A{r}="","",N(B{r})+N(C{r}))'
        ws[f"D{r}"].fill = FORMULA_FILL
    for r, (item, internal, customer) in enumerate(ORDERS, start=2):
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = item, internal or None, customer or None


# Components columns. A-H typed; I onward are formulas.
COMP_HEADERS = [
    ("Part", 20), ("Item", 22), ("Per item", 7), ("Color 1", 11), ("Color 2", 11), ("Color 3", 11),
    ("Color 4", 11), ("Notes", 34), ("Colors", 30), ("Customer needed", 9), ("Internal needed", 9),
    ("Good", 7), ("B-stock", 7), ("Customer left", 9), ("Internal left", 9), ("Left", 7),
    ("Customer sets ready", 9), ("Internal sets ready", 9), ("Fits loaded colors", 8),
    ("_planner rank", 6), ("_first of color set", 6), ("_color set #", 6),
]


def components_formulas(r):
    loaded = f"'Plate planner'!$A${LOADED_FIRST}:$A${LOADED_LAST}"
    fits = ",".join(f'OR({c}{r}="",COUNTIF({loaded},{c}{r})>0)' for c in "DEFG")
    order = lambda col: f"IFERROR(INDEX(Orders!${col}:${col},MATCH($B{r},Orders!$A:$A,0)),0)"  # noqa: E731
    blank = f'$A{r}=""'
    return {
        "I": f'=IF({blank},"",D{r}&IF(E{r}="","",", "&E{r})&IF(F{r}="","",", "&F{r})&IF(G{r}="","",", "&G{r}))',
        "J": f'=IF({blank},"",$C{r}*N({order("C")}))',
        "K": f'=IF({blank},"",$C{r}*N({order("B")}))',
        "L": f"=IF({blank},\"\",SUMIF('Print log'!$C:$C,$A{r},'Print log'!$D:$D))",
        "M": f"=IF({blank},\"\",SUMIF('Print log'!$C:$C,$A{r},'Print log'!$E:$E))",
        # Customer orders take good prints first.
        "N": f'=IF({blank},"",MAX(0,J{r}-L{r}))',
        # Internal orders take B-stock plus any good prints beyond the customer order.
        "O": f'=IF({blank},"",MAX(0,K{r}-M{r}-MAX(0,L{r}-J{r})))',
        "P": f'=IF({blank},"",N{r}+O{r})',
        "Q": f'=IF({blank},"",INT(MIN(L{r},J{r})/$C{r}))',
        "R": f'=IF({blank},"",INT((K{r}-O{r})/$C{r}))',
        "S": f'=IF({blank},"",IF(AND(COUNTA({loaded})>0,P{r}>0,{fits}),1,0))',
        # Most-needed first; ties keep sheet order.
        "T": f'=IF(S{r}=1,COUNTIFS($S$2:$S${COMPONENT_ROWS + 1},1,$P$2:$P${COMPONENT_ROWS + 1},">"&P{r})'
        f'+COUNTIFS($S$2:$S{r},1,$P$2:$P{r},P{r}),"")',
        "U": f'=IF(OR({blank},N(P{r})=0),0,IF(COUNTIFS($I$2:$I{r},$I{r},$P$2:$P{r},">0")=1,1,0))',
        "V": f'=IF(U{r}=1,SUM($U$2:$U{r}),"")',
    }


def build_components(wb):
    ws = wb.create_sheet("Components")
    header(ws, [h for h, _ in COMP_HEADERS], [w for _, w in COMP_HEADERS])
    for r in range(2, COMPONENT_ROWS + 2):
        for col, formula in components_formulas(r).items():
            ws[f"{col}{r}"] = formula
            ws[f"{col}{r}"].fill = FORMULA_FILL
    for r, (part, item, per, colors, notes) in enumerate(COMPONENTS, start=2):
        colors = (colors + [None] * 4)[:4]
        for c, v in enumerate([part, item, per, *colors, notes or None], start=1):
            ws.cell(row=r, column=c, value=v)
    for col in "TUV":
        ws.column_dimensions[col].hidden = True
    ws.column_dimensions["H"].width = 34
    for r in range(2, COMPONENT_ROWS + 2):
        ws[f"H{r}"].alignment = Alignment(wrap_text=True, vertical="top")

    last = COMPONENT_ROWS + 1
    dropdown(ws, f"B2:B{last}", f"Orders!$A$2:$A${PROGRESS_ROWS + 1}")
    dropdown(ws, f"D2:G{last}", f"Lists!$A$2:$A$60")
    ws.conditional_formatting.add(
        f"A2:S{last}", FormulaRule(formula=['AND($A2<>"",$P2=0)'], fill=PatternFill("solid", fgColor="D9EAD3"))
    )
    ws.conditional_formatting.add(
        f"A2:S{last}", FormulaRule(formula=['AND($A2<>"",$S2=1)'], fill=PatternFill("solid", fgColor="FCE8B2"))
    )


def build_log(wb):
    ws = wb.create_sheet("Print log")
    header(ws, ["Date", "Who", "Part", "Good", "B-stock", "Printer", "Notes"], [11, 10, 22, 7, 8, 22, 36])
    for r, row in enumerate(PRINT_LOG, start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v or None)
    for r in range(2, LOG_ROWS + 1):
        ws[f"A{r}"].number_format = "yyyy-mm-dd"
    dropdown(ws, f"C2:C{LOG_ROWS}", f"Components!$A$2:$A${COMPONENT_ROWS + 1}")
    dropdown(ws, f"F2:F{LOG_ROWS}", "Printers!$A$2:$A$60")


def build_progress(wb):
    ws = wb.create_sheet("Progress")
    labels = ["Item", "Customer needed", "Customer ready", "Internal needed", "Internal ready", "Pieces left", "Status"]
    header(ws, labels, [26, 10, 10, 10, 10, 9, 12])
    ws.freeze_panes = "A2"
    for i in range(PROGRESS_ROWS):
        r, o = i + 2, i + 2  # Progress row r mirrors Orders row o
        blank = f'$A{r}=""'
        has_parts = f"COUNTIF(Components!$B:$B,$A{r})>0"
        ready = lambda col, need: (  # noqa: E731
            f'=IF({blank},"",IF({has_parts},MIN({need}{r},_xlfn.MINIFS(Components!${col}:${col},Components!$B:$B,$A{r})),"no parts"))'
        )
        printed = f"SUMIF(Components!$B:$B,$A{r},Components!$L:$L)+SUMIF(Components!$B:$B,$A{r},Components!$M:$M)"
        cells = {
            "A": f'=IF(Orders!A{o}="","",Orders!A{o})',
            "B": f'=IF({blank},"",N(Orders!C{o}))',
            "C": ready("Q", "B"),
            "D": f'=IF({blank},"",N(Orders!B{o}))',
            "E": ready("R", "D"),
            "F": f'=IF({blank},"",SUMIF(Components!$B:$B,$A{r},Components!$P:$P))',
            "G": f'=IF({blank},"",IF(B{r}+D{r}=0,"No order",IF(F{r}=0,"Done",IF({printed}>0,"In progress","Not started"))))',
        }
        for col, formula in cells.items():
            ws[f"{col}{r}"] = formula
            ws[f"{col}{r}"].fill = FORMULA_FILL
    last = PROGRESS_ROWS + 1
    for status, color in (("Done", "B6D7A8"), ("In progress", "FCE8B2"), ("No order", "EEEEEE")):
        ws.conditional_formatting.add(f"G2:G{last}", FormulaRule(formula=[f'$G2="{status}"'], fill=PatternFill("solid", fgColor=color)))

    ws["I1"], ws["I2"] = "Keyboard switches needed", "=SUM(Orders!D:D)"
    ws["I1"].font = Font(bold=True)
    ws["I3"] = "Sets ready = copies of the item's slowest part."
    ws["I4"] = "B-stock only fills internal orders."
    ws["I3"].font = ws["I4"].font = NOTE_FONT
    ws.column_dimensions["I"].width = 44


def build_planner(wb):
    ws = wb.create_sheet("Plate planner")
    ws["A1"] = "Plate planner"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (
        "Type the colors loaded in your AMS or toolheads in A5:A12 (one per cell). Left: every part still "
        "needed that prints with only those colors, most-needed first. Right: every color set still needed."
    )
    ws["A2"].font = NOTE_FONT
    ws["A4"] = "Loaded colors"
    for r in range(LOADED_FIRST, LOADED_LAST + 1):
        ws[f"A{r}"].fill = INPUT_FILL
    dropdown(ws, f"A{LOADED_FIRST}:A{LOADED_LAST}", "Lists!$A$2:$A$60")
    ws[f"A{LOADED_FIRST}"], ws[f"A{LOADED_FIRST + 1}"] = "white", "black"  # example; replace with what's loaded

    fits_cols = [("C", "Part", "A", 20), ("D", "Item", "B", 22), ("E", "Colors", "I", 30),
                 ("F", "Customer left", "N", 9), ("G", "Left", "P", 7)]
    sets_cols = [("I", "Color set still needed", 30), ("J", "Pieces left", 9), ("K", "Parts", 7)]
    for col, label, *_ in fits_cols + sets_cols:
        ws[f"{col}4"] = label
        ws[f"{col}4"].fill = HEADER_FILL
        ws[f"{col}4"].font = HEADER_FONT
    for col, _, _, width in fits_cols:
        ws.column_dimensions[col].width = width
    for col, _, width in sets_cols:
        ws.column_dimensions[col].width = width
    ws.column_dimensions["A"].width = 16

    for k in range(1, PLANNER_ROWS + 1):
        r = k + 4
        match = f"MATCH({k},Components!$T:$T,0)"
        for col, _, src, _ in fits_cols:
            ws[f"{col}{r}"] = f'=IFERROR(INDEX(Components!${src}:${src},{match}),"")'
        ws[f"I{r}"] = f'=IFERROR(INDEX(Components!$I:$I,MATCH({k},Components!$V:$V,0)),"")'
        ws[f"J{r}"] = f'=IF(I{r}="","",SUMIF(Components!$I:$I,I{r},Components!$P:$P))'
        ws[f"K{r}"] = f'=IF(I{r}="","",COUNTIFS(Components!$I:$I,I{r},Components!$P:$P,">0"))'
        for col in "CDEFGIJK":
            ws[f"{col}{r}"].fill = FORMULA_FILL


def build_printers(wb):
    ws = wb.create_sheet("Printers")
    header(ws, ["Printer", "With", "Colors at once", "Notes"], [28, 12, 9, 24])
    for r, row in enumerate(PRINTERS, start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v or None)


def build_how_to(wb):
    ws = wb.create_sheet("How to use")
    ws.column_dimensions["A"].width = 110
    lines = [
        ("Clicker print tracker", TITLE_FONT),
        ("After printing", Font(bold=True, color=RED)),
        ("Add a line to Print log: date, who, part, how many good, how many B-stock. That's the only place counts go.", None),
        ("Components, Progress and the Plate planner update from it.", None),
        ("Setting up items", Font(bold=True, color=RED)),
        ("Orders: one row per item with internal and customer quantities. Total = keyboard switches needed.", None),
        ("Components: one row per printed piece. Colors on one row print together as one multicolor print.", None),
        ("If an item prints as separate pieces, give each piece its own row with the same Item.", None),
        ("Yellow-highlighted Components rows fit the colors loaded on the Plate planner. Green rows are finished.", None),
        ("Blue cells are formulas. Type only in the white columns.", None),
        ("Rules", Font(bold=True, color=RED)),
        ("Customer orders take good prints first. B-stock, plus any good prints beyond the customer order, fill internal orders.", None),
        ("Sets ready for an item = copies of its slowest part.", None),
        ("A part shows on the Plate planner only when every one of its colors is loaded.", None),
        ("Use the same spelling for a color everywhere (pick from the dropdowns; add new colors on Lists).", None),
        ("About", Font(bold=True, color=RED)),
        (f"Version {VERSION}, {DATE}, {AUTHOR}. Built by tools/print_tracker/build_clicker_tracker_v{VERSION}.py in the Rambots repo.", None),
        (f"Changelog: v1.0 ({DATE}) Initial release.", None),
    ]
    for r, (text, font) in enumerate(lines, start=1):
        cell = ws.cell(row=r, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True)
        if font:
            cell.font = font


def build_lists(wb):
    ws = wb.create_sheet("Lists")
    header(ws, ["Colors"], [16])
    ws.freeze_panes = None
    for r, color in enumerate(COLORS, start=2):
        ws[f"A{r}"] = color


def build():
    wb = Workbook()
    wb.remove(wb.active)
    build_progress(wb)
    build_components(wb)
    build_log(wb)
    build_planner(wb)
    build_orders(wb)
    build_printers(wb)
    build_how_to(wb)
    build_lists(wb)
    return wb


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--out", default=f"clicker_print_tracker_v{VERSION}.xlsx")
    args = parser.parse_args()
    build().save(args.out)
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()

# CHANGELOG
# v1.0 (2026-10-20): Initial release. Orders, Components, Print log, Progress, Plate planner, Printers,
#   How to use and Lists tabs, seeded from the "Robotics internal clicker prints" Google Sheet.
