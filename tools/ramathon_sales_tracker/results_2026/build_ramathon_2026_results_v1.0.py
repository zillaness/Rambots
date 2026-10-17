"""
file: build_ramathon_2026_results_v1.0.py
version: 1.0
author: Sam Cao
created: 2026-10-17
last_updated: 2026-10-17
description: Builds the Ram-a-thon 2026 results workbook (formula-driven) and a metrics JSON from the sales laptop's CSV export.
ai_update: Update last_updated and version. Rename file to match. Append changelog at bottom.

Usage:
    python3 build_ramathon_2026_results_v1.0.py <sales.csv> <out.xlsx> <out_metrics.json> [<summary.csv>]

Everything the workbook shows is a formula over the "Raw Data" tab plus the blue
input cells, so correcting a stock count or an estimate recalculates the rest.
The metrics JSON repeats the same math in Python for the findings write-up and
the infographic, and is used to check the recalculated workbook.
"""

import csv
import itertools
import json
import sys
from collections import Counter, defaultdict
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Inputs from Sam (2026-10-17 chat) and the tracker README (built-in stock)
# ---------------------------------------------------------------------------

# Oct 16 counts built into the tracker. Egg tart was restocked the morning of
# the event, so its true starting count is unknown (at least 11).
STARTING_STOCK = {
    "Egg tart": None, "Hot dog bun": 22, "S'mores": 9, "Chocolate bar": 15,
    "Avocado": 12, "Dumpling": 24, "Creeper": 10, "Oreo": 4,
    "Spam musubi": 6, "Salmon nigiri": 23, "LEGO brick": 4, "WD-40": 15,
    "Creeper infinity cube": 9, "Burger": None,
}
NOT_RANKED = {
    "Oreo": "Not ranked: probably never put out on the table (Sam).",
    "Burger": "Not ranked: not in the stock count, 1 sold.",
}
THEMES = {
    "Salmon nigiri": "Asian food", "Dumpling": "Asian food",
    "Spam musubi": "Asian food", "Egg tart": "Asian food",
    "Hot dog bun": "Snacks & treats", "Burger": "Snacks & treats",
    "S'mores": "Snacks & treats", "Chocolate bar": "Snacks & treats",
    "Oreo": "Snacks & treats", "Avocado": "Snacks & treats",
    "Creeper": "Minecraft", "Creeper infinity cube": "Minecraft",
    "WD-40": "Workshop", "LEGO brick": "Workshop",
    "Full kit": "Copters", "Kit w/ toroidal": "Copters",
}
THEME_ORDER = ["Asian food", "Snacks & treats", "Minecraft", "Workshop", "Copters"]
CLICKER_THEMES = THEME_ORDER[:4]

EVENT_INPUTS = {
    "counted": 830,          # tickets counted at the end
    "gross": 900,            # Sam: ~5 tix per helper for food -> close to 900
    "cc_low": 500,           # clickers + copters incl. untracked blowout
    "cc_high": 530,
    "battle_low": 50,
    "battle_high": 60,
    "plinko_tix_per_drop": 1,
    "battle_tix_per_match": 5,
    "cca_cut": 0.10,
    "dollars_per_ticket": 1,
    "copters_last_year_net": 450,
}
GOALS = {"Clickers": 500, "Copters": 500, "Balloon Battle": 200, "Plinko": 100}
COPTER_PARTS = {"Launchers": 93, "Propellers": 158, "Toroidal propellers": 44, "Ripcords": 105}

# Sam: units sold mean less than rate because we made unequal amounts of each design.
RAW_KEEP = ["record", "opened_at", "closed_at", "category", "item", "qty", "item_tickets",
            "record_total_tickets", "age", "gender", "bonus_cards"]

RATE_WEIGHT = 0.7
CAP_PER_SALE = 3
BUFFER = 0.20

# ---------------------------------------------------------------------------
# Load
# ---------------------------------------------------------------------------

def parse_dt(s):
    return datetime.strptime(s, "%Y-%m-%d %H:%M:%S")


def load(path):
    with open(path, newline="") as f:
        reader = csv.DictReader(f)
        header = reader.fieldnames
        rows = list(reader)
    for r in rows:
        r["_opened"] = parse_dt(r["opened_at"])
        r["_closed"] = parse_dt(r["closed_at"])
        r["_qty"] = int(r["qty"])
        r["_tix"] = int(r["item_tickets"])
    return header, rows


# ---------------------------------------------------------------------------
# Python-side metrics (mirror the workbook formulas)
# ---------------------------------------------------------------------------

def hours(td):
    return td.total_seconds() / 3600


def compute(rows):
    m = {}
    sales = {}
    for r in rows:
        sales.setdefault(r["record"], []).append(r)
    win_start = min(r["_opened"] for r in rows)
    win_end = max(r["_closed"] for r in rows)
    win_h = hours(win_end - win_start)
    m["window"] = {"start": win_start.strftime("%H:%M:%S"), "end": win_end.strftime("%H:%M:%S"), "hours": win_h}

    clicker_rows = [r for r in rows if r["category"] != "Copters"]
    copter_rows = [r for r in rows if r["category"] == "Copters"]
    m["tracked"] = {
        "sales": len(sales),
        "rows": len(rows),
        "tickets": sum(r["_tix"] for r in rows),
        "clicker_tickets": sum(r["_tix"] for r in clicker_rows),
        "clicker_units": sum(r["_qty"] for r in clicker_rows),
        "copter_tickets": sum(r["_tix"] for r in copter_rows),
        "copter_kits": sum(r["_qty"] for r in copter_rows),
        "free_cards": sum(int(v[0]["bonus_cards"] or 0) for v in sales.values()),
    }
    avg_tix_per_clicker = m["tracked"]["clicker_tickets"] / m["tracked"]["clicker_units"]
    m["avg_tix_per_clicker"] = avg_tix_per_clicker

    # Per-design
    designs = {}
    for d in STARTING_STOCK:
        drs = sorted([r for r in rows if r["item"] == d], key=lambda r: r["_closed"])
        units = sum(r["_qty"] for r in drs)
        stock = STARTING_STOCK[d]
        sold_out = stock is not None and units >= stock
        sold_out_at = drs[-1]["_closed"] if (sold_out and drs) else None
        in_stock_h = hours((sold_out_at or win_end) - win_start)
        units_for_rate = sum(min(r["_qty"], CAP_PER_SALE) for r in drs)
        rate = units_for_rate / in_stock_h if in_stock_h > 0 else 0
        hours_out = hours(win_end - sold_out_at) if sold_out_at else 0
        lost_units = rate * hours_out
        designs[d] = {
            "theme": THEMES[d], "stock": stock, "units": units,
            "tickets": sum(r["_tix"] for r in drs),
            "sell_through": (units / stock) if stock else None,
            "sold_out": sold_out,
            "sold_out_at": sold_out_at.strftime("%H:%M") if sold_out_at else None,
            "first_sale": drs[0]["_closed"].strftime("%H:%M") if drs else None,
            "in_stock_hours": in_stock_h, "units_for_rate": units_for_rate,
            "rate": rate, "hours_out": hours_out, "lost_units": lost_units,
            "lost_tickets": lost_units * avg_tix_per_clicker,
            "left_over": (max(0, stock - units) if stock is not None else None),
            "ideal_next_year": -(-rate * win_h * (1 + BUFFER) // 1),
            "ranked": d not in NOT_RANKED,
        }
    ranked = [d for d in designs if designs[d]["ranked"]]
    max_rate = max(designs[d]["rate"] for d in ranked)
    max_units = max(designs[d]["units"] for d in ranked)
    for d in ranked:
        x = designs[d]
        x["rate_score"] = x["rate"] / max_rate
        x["volume_score"] = x["units"] / max_units
        x["blended"] = 100 * (RATE_WEIGHT * x["rate_score"] + (1 - RATE_WEIGHT) * x["volume_score"])
    for key, name in (("units", "units_rank"), ("rate", "rate_rank"), ("blended", "blended_rank")):
        vals = sorted((designs[d][key] for d in ranked), reverse=True)
        for d in ranked:
            designs[d][name] = vals.index(designs[d][key]) + 1
    m["designs"] = designs
    # Primary ranking is sales rate (Sam, Oct 17); blended order kept for reference.
    m["ranking_order"] = sorted(ranked, key=lambda d: -designs[d]["rate"])
    m["blended_order"] = sorted(ranked, key=lambda d: -designs[d]["blended"])
    m["lost"] = {
        "units": sum(designs[d]["lost_units"] for d in ranked),
        "tickets": sum(designs[d]["lost_tickets"] for d in ranked),
        "left_over_units": sum(designs[d]["left_over"] or 0 for d in designs),
    }

    # Booth pace before / after the last of the big three sold out
    # All three sold out, so their last sale times are their sell-out times.
    cutoff = max(r["_closed"] for r in rows if r["item"] in ("Creeper", "Dumpling", "Avocado"))
    before = [r for r in clicker_rows if r["_closed"] <= cutoff]
    after = [r for r in clicker_rows if r["_closed"] > cutoff]
    m["pace"] = {
        "cutoff": cutoff.strftime("%H:%M:%S"),
        "before_units_per_hour": sum(r["_qty"] for r in before) / hours(cutoff - win_start),
        "after_units_per_hour": sum(r["_qty"] for r in after) / hours(win_end - cutoff),
        "sales_before": len({r["record"] for r in before}),
        "sales_after": len({r["record"] for r in after}),
    }

    # Baskets
    baskets = []
    for rec, v in sales.items():
        items = Counter()
        for r in v:
            items[r["item"]] += r["_qty"]
        baskets.append({
            "record": rec, "opened": v[0]["_opened"], "closed": v[0]["_closed"],
            "age": v[0]["age"], "gender": v[0]["gender"],
            "tickets": int(v[0]["record_total_tickets"]),
            "units": sum(items.values()), "items": items,
            "themes": Counter(THEMES[i] for i in items for _ in range(items[i])),
            "free_cards": int(v[0]["bonus_cards"] or 0),
        })
    m["baskets"] = {
        "count": len(baskets),
        "avg_tickets": sum(b["tickets"] for b in baskets) / len(baskets),
        "tickets_hist": dict(sorted(Counter(b["tickets"] for b in baskets).items())),
        "units_hist": dict(sorted(Counter(b["units"] for b in baskets).items())),
        "multi_design": sum(1 for b in baskets if len(b["items"]) > 1),
        "multi_theme": sum(1 for b in baskets if len(b["themes"]) > 1),
        "free_card_sales": sum(1 for b in baskets if b["free_cards"] > 0),
        "biggest": max(baskets, key=lambda b: b["tickets"])["record"],
        "checkout_seconds_median": sorted((b["closed"] - b["opened"]).total_seconds() for b in baskets)[len(baskets) // 2],
        "checkout_seconds_min": min((b["closed"] - b["opened"]).total_seconds() for b in baskets),
    }
    multi_unit = [b for b in baskets if b["units"] > 1]
    m["baskets"]["multi_unit"] = len(multi_unit)
    m["baskets"]["same_design_multi"] = sum(1 for b in multi_unit if len(b["items"]) == 1)
    mixers, partners = Counter(), defaultdict(set)
    for b in baskets:
        if len(b["items"]) > 1:
            for i in b["items"]:
                mixers[i] += 1
                partners[i] |= set(b["items"]) - {i}
    m["mixers"] = [[i, n, len(partners[i])] for i, n in mixers.most_common()]
    pairs = Counter()
    for b in baskets:
        for a, c in itertools.combinations(sorted(b["items"]), 2):
            pairs[(a, c)] += 1
    m["pairs"] = [[a, c, n] for (a, c), n in pairs.most_common()]
    tpairs = Counter()
    for b in baskets:
        ts = sorted(b["themes"])
        if len(b["items"]) > 1 and len(ts) == 1:
            tpairs[(ts[0], ts[0])] += 1
        for a, c in itertools.combinations(ts, 2):
            tpairs[(a, c)] += 1
    m["theme_pairs"] = [[a, c, n] for (a, c), n in tpairs.most_common()]

    # Themes
    th = {}
    for t in THEME_ORDER:
        ds = [d for d in designs if designs[d]["theme"] == t]
        known = [d for d in ds if designs[d]["stock"] is not None and d != "Oreo"]
        th[t] = {
            "designs": ds,
            "units": sum(designs[d]["units"] for d in ds) if t != "Copters" else m["tracked"]["copter_kits"],
            "stock_known": sum(designs[d]["stock"] for d in known),
            "units_known": sum(designs[d]["units"] for d in known),
        }
        th[t]["sell_through_known"] = (th[t]["units_known"] / th[t]["stock_known"]) if th[t]["stock_known"] else None
    m["themes"] = th

    # Demographics (units, by who the items were for)
    demo = defaultdict(Counter)
    demo_theme = defaultdict(Counter)
    for r in rows:
        for seg in (r["age"] or "Untagged age", r["gender"] or "Untagged gender"):
            demo[seg][r["item"]] += r["_qty"]
            demo_theme[seg][THEMES[r["item"]]] += r["_qty"]
    m["demographics"] = {
        "sales_by_age": dict(Counter(b["age"] or "Untagged" for b in baskets)),
        "sales_by_gender": dict(Counter(b["gender"] or "Untagged" for b in baskets)),
        "units_by_segment": {k: dict(v.most_common()) for k, v in demo.items()},
        "theme_units_by_segment": {k: dict(v.most_common()) for k, v in demo_theme.items()},
    }
    groups = [
        ("Kids", "Kid", None), ("Teens", "Teen", None), ("Adults", "Adult", None),
        ("Boys and men", None, "Male"), ("Girls and women", None, "Female"),
        ("Kid boys", "Kid", "Male"), ("Kid girls", "Kid", "Female"),
        ("Teen boys", "Teen", "Male"), ("Teen girls", "Teen", "Female"),
        ("Adult men", "Adult", "Male"), ("Adult women", "Adult", "Female"),
    ]
    top = []
    for label, age, gen in groups:
        c = Counter()
        for r in rows:
            if r["category"] == "Copters":
                continue
            # Count customers (sales that picked the design), not units, so one
            # bulk buyer doesn't crown a favorite. Each row is one design in one sale.
            if (age is None or r["age"] == age) and (gen is None or r["gender"] == gen):
                c[r["item"]] += 1
        tiers = defaultdict(list)
        for item, q in c.items():
            tiers[q].append(item)
        parts, listed = [], 0
        for q in sorted(tiers, reverse=True):
            if listed >= 3:
                break
            names = sorted(tiers[q])
            if listed and len(names) > 3:
                parts.append(f"{len(names)} others ({q} each)")
            else:
                parts.append(", ".join(names) + (f" ({q} each)" if len(names) > 1 else f" ({q})"))
            listed += len(names)
        top.append({"group": label, "age": age, "gender": gen,
                    "top": "; ".join(parts), "top_list": [[i, q] for i, q in c.most_common()]})
    m["top_by_group"] = top
    male_share = {}
    for d in designs:
        mm = demo["Male"][d]
        ff = demo["Female"][d]
        if mm + ff:
            male_share[d] = {"male": mm, "female": ff, "male_share": mm / (mm + ff)}
    m["male_share_by_design"] = male_share

    # Timeline (half-hour buckets)
    t0 = datetime(2026, 10, 17, 10, 0)
    tl = []
    for i in range(8):
        a, b = t0 + timedelta(minutes=30 * i), t0 + timedelta(minutes=30 * (i + 1))
        bs = [x for x in baskets if a <= x["closed"] < b]
        tl.append({"start": a.strftime("%H:%M"), "sales": len(bs),
                   "tickets": sum(x["tickets"] for x in bs), "units": sum(x["units"] for x in bs)})
    m["timeline"] = tl

    # Event estimates
    e = EVENT_INPUTS
    cc_mid = (e["cc_low"] + e["cc_high"]) / 2
    battle_mid = (e["battle_low"] + e["battle_high"]) / 2
    plinko_low = e["gross"] - e["cc_high"] - e["battle_high"]
    plinko_high = e["gross"] - e["cc_low"] - e["battle_low"]
    plinko_mid = e["gross"] - cc_mid - battle_mid
    copters = m["tracked"]["copter_tickets"]
    m["event"] = {
        "gross": e["gross"], "counted": e["counted"], "helper_food": e["gross"] - e["counted"],
        "net_after_cca": e["gross"] * (1 - e["cca_cut"]) * e["dollars_per_ticket"],
        "clickers_mid": cc_mid - copters, "copters": copters,
        "battle_mid": battle_mid, "battle_matches_low": e["battle_low"] / e["battle_tix_per_match"],
        "battle_matches_high": e["battle_high"] / e["battle_tix_per_match"],
        "plinko_low": plinko_low, "plinko_high": plinko_high, "plinko_mid": plinko_mid,
        "plinko_drops_mid": plinko_mid / e["plinko_tix_per_drop"],
        "goal_total": sum(GOALS.values()),
        "copters_this_year_net": copters * (1 - e["cca_cut"]),
        "copters_last_year_net": e["copters_last_year_net"],
    }
    m["event"]["copter_change"] = m["event"]["copters_this_year_net"] / e["copters_last_year_net"] - 1
    kits = Counter()
    for r in copter_rows:
        kits[r["item"]] += r["_qty"]
    m["copters"] = {
        "kits": dict(kits),
        "parts_left": {
            "Launchers": COPTER_PARTS["Launchers"] - sum(kits.values()),
            "Propellers": COPTER_PARTS["Propellers"] - kits["Full kit"],
            "Toroidal propellers": COPTER_PARTS["Toroidal propellers"] - kits["Kit w/ toroidal"],
            "Ripcords": COPTER_PARTS["Ripcords"] - sum(kits.values()),
        },
    }
    return m, baskets


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

RED = "A31F36"
YELLOW = "F4A81C"
FONT = "Arial"
F_BODY = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True, color=RED)
F_SUB = Font(name=FONT, size=10, italic=True, color="555555")
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
FILL_HEAD = PatternFill("solid", fgColor=RED)
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_BAND = PatternFill("solid", fgColor="F7F7F7")
FILL_TOTAL = PatternFill("solid", fgColor="FBE9C6")
THIN = Side(style="thin", color="D0D0D0")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
TIME_FMT = "h:mm AM/PM"
DT_FMT = "yyyy-mm-dd hh:mm:ss"


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_SUB


def header(ws, row, labels, col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX


def put(ws, row, col, value, fmt=None, font=None, fill=None, note=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font or F_BODY
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if note:
        c.comment = Comment(note, "Sam Cao")
    return c


def inp(ws, row, col, value, fmt=None, note=None):
    return put(ws, row, col, value, fmt=fmt, font=F_INPUT, fill=FILL_INPUT, note=note)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build(header_cols, rows, m, baskets, out_path):
    wb = Workbook()
    n = len(rows)
    last = n + 1  # last Raw Data row
    RD = "'Raw Data'"

    def rd(col):
        return f"{RD}!${col}$2:${col}${last}"

    # Column letters in Raw Data
    cidx = {h: get_column_letter(i + 1) for i, h in enumerate(header_cols)}
    C_ITEM, C_QTY, C_TIX = cidx["item"], cidx["qty"], cidx["item_tickets"]
    C_CAT, C_REC, C_AGE, C_GEN = cidx["category"], cidx["record"], cidx["age"], cidx["gender"]
    C_OPEN, C_CLOSE, C_BONUS = cidx["opened_at"], cidx["closed_at"], cidx["bonus_cards"]
    C_THEME = get_column_letter(len(header_cols) + 1)
    C_CAPQ = get_column_letter(len(header_cols) + 2)

    designs = m["designs"]
    order = m["ranking_order"] + [d for d in designs if not designs[d]["ranked"]]
    n_ranked = len(m["ranking_order"])

    # ------------------------------------------------------------------ Read Me
    ws = wb.active
    ws.title = "Read Me"
    title(ws, "Ram-a-thon 2026 results (Rambots, Team 2204)",
          "CCA Hayward, Sat Oct 17 2026. Built from the sales laptop's CSV export (Laptop 1, exported 1:51 PM).")
    lines = [
        ("What's tracked", "Every clicker and copter sale rung up on the sales laptop: 69 sales, 10:01 AM to 1:20 PM."),
        ("What's not tracked", "Plinko (bucket not counted), Balloon Battle (laptop not exported), the Prize Table, and the end-of-day blowout sale. Those numbers are estimates and live in blue input cells."),
        ("Blue cells", "Inputs. Change them (a stock count, an estimate, a weight) and every formula updates."),
        ("Leaderboard", "Designs ranked by sales rate, plus the most popular designs with each age and gender group."),
        ("Design Ranking", "The math behind the leaderboard. Sales rate is the main rank because we made unequal amounts of each design. The blended rank (70% rate, 30% units) is there for reference; the weights are inputs."),
        ("Sales rate", "Units sold divided by hours the design was on the table. A sale counts at most 3 of one design toward the rate, so the parent who bought 8 Dumplings doesn't skew it."),
        ("Stock & Lost Sales", "For designs that sold out: rate x hours left after the sellout = units we could have sold with unlimited stock. Also ideal stock for next year."),
        ("Event Totals", "Tracked vs estimated tickets per booth, Plinko inferred as gross minus everything else, goals, and $ after CCA's 10%."),
        ("Themes / Pairings", "Which kinds of clickers sold, and which designs were bought together."),
        ("Demographics", "Who items were for (age and gender tags). About a third of sales are untagged, and Adults are only 4 sales, so treat small groups as hints."),
        ("Findings", "Conclusions, hypotheses, and things to change next year."),
        ("Raw Data", "The 11 export columns that hold data. The battle, prize, and laptop columns were empty or constant; the full export is ramathon_sales_laptop1_2026-10-17.csv in the repo."),
        ("Sources", "Sales CSV (tracker v1.8.5). Starting stock: Sam's Oct 16 counts built into the tracker. Ticket totals, Battle, blowout, helper tickets, Oreo, Egg tart, Plinko price, last year's copters: Sam, Oct 17 chat."),
        ("Version", "ramathon_2026_results_v1.0.xlsx, built 2026-10-17 by build_ramathon_2026_results_v1.0.py"),
    ]
    for i, (k, v) in enumerate(lines, 4):
        ws.cell(row=i, column=1, value=k).font = F_BOLD
        c = ws.cell(row=i, column=2, value=v)
        c.font = F_BODY
        c.alignment = WRAP
    widths(ws, [22, 110])

    # ------------------------------------------------- Findings, Leaderboard
    wsF = wb.create_sheet("Findings")
    wsL = wb.create_sheet("Leaderboard")
    NB = len(baskets) + 1

    # ------------------------------------------------------------ Design Ranking
    ws = wb.create_sheet("Design Ranking")
    title(ws, "Clicker design ranking", "Sorted by sales rate. Blended score = rate weight x (rate / best rate) + volume weight x (units / most units), out of 100.")
    DR = "'Design Ranking'"
    inputs = [
        ("Rate weight", RATE_WEIGHT, "0%", "Share of the blended score from sales rate. Sam: units sold mean less because we made unequal amounts, so rate gets 70%."),
        ("Volume weight", "=1-B4", "0%", None),
        ("Max units of one design per sale counted toward rate", CAP_PER_SALE, "0", "3 matches the 3-for-10 deal. Only one sale (8 Dumplings) is affected."),
        ("Selling window start", f"=MIN({rd(C_OPEN)})", TIME_FMT, None),
        ("Selling window end", f"=MAX({rd(C_CLOSE)})", TIME_FMT, None),
        ("Window hours", "=(B8-B7)*24", "0.00", None),
    ]
    for i, (lab, val, fmt, note) in enumerate(inputs, 4):
        ws.cell(row=i, column=1, value=lab).font = F_BOLD
        if isinstance(val, str) and val.startswith("="):
            put(ws, i, 2, val, fmt)
        else:
            inp(ws, i, 2, val, fmt, note)
    hdr_row = 11
    header(ws, hdr_row, ["Design", "Theme", "Starting stock", "On table from", "Units sold",
                         "Tickets rung up", "Sell-through", "Sold out?", "Sold out at",
                         "Hours in stock", "Units counted for rate", "Sales rate (units/hr)",
                         "Rate score", "Volume score", "Blended score", "Blended rank",
                         "Units rank", "Rate rank", "Notes"])
    first = hdr_row + 1
    last_ranked = first + n_ranked - 1
    design_row = {}
    for i, d in enumerate(order):
        r = first + i
        design_row[d] = r
        x = designs[d]
        put(ws, r, 1, d, font=F_BOLD)
        inp(ws, r, 2, x["theme"])
        note = None
        if d == "Egg tart":
            note = "Restocked the morning of the event after the Oct 16 count of 8; true starting count unknown (at least 11). Type it here if you find it."
        inp(ws, r, 3, x["stock"], "0", note)
        tnote = ("First Egg tart sale was 11:07 AM, about 45 min after most designs started selling. If it only reached the table then, set this to 11:00 AM and its rate goes up."
                 if d == "Egg tart" else None)
        inp(ws, r, 4, "=$B$7", TIME_FMT, tnote)
        put(ws, r, 5, f"=SUMIFS({rd(C_QTY)},{rd(C_ITEM)},A{r})", "0")
        put(ws, r, 6, f"=SUMIFS({rd(C_TIX)},{rd(C_ITEM)},A{r})", "0")
        put(ws, r, 7, f'=IF(C{r}="","",E{r}/C{r})', "0%")
        put(ws, r, 8, f'=IF(C{r}="","Unknown",IF(E{r}>=C{r},"Yes","No"))')
        put(ws, r, 9, f'=IF(H{r}="Yes",SUMPRODUCT(MAX(({rd(C_ITEM)}=A{r})*{rd(C_CLOSE)})),"")', TIME_FMT)
        put(ws, r, 10, f'=IF(H{r}="Yes",(I{r}-D{r})*24,($B$8-D{r})*24)', "0.00")
        put(ws, r, 11, f"=SUMIFS({rd(C_CAPQ)},{rd(C_ITEM)},A{r})", "0")
        put(ws, r, 12, f"=IF(J{r}>0,K{r}/J{r},0)", "0.0")
        if designs[d]["ranked"]:
            rng = lambda col: f"${col}${first}:${col}${last_ranked}"
            put(ws, r, 13, f"=L{r}/MAX({rng('L')})", "0.00")
            put(ws, r, 14, f"=E{r}/MAX({rng('E')})", "0.00")
            put(ws, r, 15, f"=100*($B$4*M{r}+$B$5*N{r})", "0")
            put(ws, r, 16, f"=RANK(O{r},{rng('O')})", "0")
            put(ws, r, 17, f"=RANK(E{r},{rng('E')})", "0")
            put(ws, r, 18, f"=RANK(L{r},{rng('L')})", "0", font=F_BOLD, fill=FILL_TOTAL)
            notes = {
                "Creeper": "Sold out first (11:48 AM), no bulk buys.",
                "Dumpling": "One parent bought the last 8 in one sale at 12:42 PM.",
                "Avocado": "Last 3 went in one 3-for-10 sale at 12:43 PM.",
                "Salmon nigiri": "Last 4 went at 1:20 PM: sold out right at the end.",
                "Spam musubi": "Only 6 made; last 2 went at 1:20 PM.",
                "Egg tart": "Restocked day-of; sell-through unknown.",
                "LEGO brick": "Both sales at 10:41 to 10:42 AM, then none. Worth checking if it got buried.",
                "Hot dog bun": "Most units left over (12 of 22).",
            }
            put(ws, r, 19, notes.get(d, ""))
        else:
            for col in range(13, 19):
                put(ws, r, col, "")
            put(ws, r, 19, NOT_RANKED[d])
    ws.freeze_panes = ws.cell(row=first, column=2)
    widths(ws, [22, 16, 10, 11, 9, 10, 10, 9, 11, 9, 11, 11, 8, 8, 9, 9, 8, 8, 60])
    ws.row_dimensions[hdr_row].height = 42

    # --------------------------------------------------------- Stock & Lost Sales
    ws = wb.create_sheet("Stock & Lost Sales")
    title(ws, "What we ran short on, and what it cost",
          "Lost units = sales rate x hours the design was gone before the last sale of the day. Ideal stock = rate x window x (1 + buffer).")
    ws["A4"] = "Buffer for next year's stock"
    ws["A4"].font = F_BOLD
    inp(ws, 4, 2, BUFFER, "0%", "Extra stock on top of this year's demand rate. 20% is a starting guess.")
    ws["A5"] = "Avg tickets per clicker (this year)"
    ws["A5"].font = F_BOLD
    put(ws, 5, 2, f'=SUMIFS({rd(C_TIX)},{rd(C_CAT)},"Clickers & Fidget Toys")/SUMIFS({rd(C_QTY)},{rd(C_CAT)},"Clickers & Fidget Toys")', "0.00")
    ws["A6"] = "Window hours"
    ws["A6"].font = F_BOLD
    put(ws, 6, 2, f"={DR}!B9", "0.00", font=F_LINK)
    h = 8
    header(ws, h, ["Design", "Starting stock", "Units sold", "Left over", "Sold out at",
                   "Hours sold out", "Sales rate (units/hr)", "Lost units", "Lost tickets",
                   "Ideal stock next year", "Change vs this year"])
    for i, d in enumerate(m["ranking_order"]):
        r = h + 1 + i
        s = design_row[d]
        put(ws, r, 1, f"={DR}!A{s}", font=F_LINK)
        put(ws, r, 2, f'=IF({DR}!C{s}="","",{DR}!C{s})', "0", font=F_LINK)
        put(ws, r, 3, f"={DR}!E{s}", "0", font=F_LINK)
        put(ws, r, 4, f'=IF(B{r}="","",MAX(0,B{r}-C{r}))', "0")
        put(ws, r, 5, f'={DR}!I{s}', TIME_FMT, font=F_LINK)
        put(ws, r, 6, f'=IF({DR}!H{s}="Yes",({DR}!$B$8-{DR}!I{s})*24,0)', "0.00")
        put(ws, r, 7, f"={DR}!L{s}", "0.0", font=F_LINK)
        put(ws, r, 8, f"=G{r}*F{r}", "0.0")
        put(ws, r, 9, f"=H{r}*$B$5", "0")
        put(ws, r, 10, f"=ROUNDUP(G{r}*$B$6*(1+$B$4),0)", "0", fill=FILL_TOTAL)
        put(ws, r, 11, f'=IF(B{r}="","",J{r}-B{r})', '+0;-0;0')
    tr = h + 1 + n_ranked
    put(ws, tr, 1, "Total", font=F_BOLD, fill=FILL_TOTAL)
    for col, fmt in ((2, "0"), (3, "0"), (4, "0"), (8, "0.0"), (9, "0"), (10, "0")):
        L = get_column_letter(col)
        put(ws, tr, col, f"=SUM({L}{h+1}:{L}{tr-1})", fmt, font=F_BOLD, fill=FILL_TOTAL)
    for col in (5, 6, 7, 11):
        put(ws, tr, col, "", fill=FILL_TOTAL)
    put(ws, tr + 1, 1, "Oreo (not put out)", font=F_BODY)
    put(ws, tr + 1, 2, f"={DR}!C{design_row['Oreo']}", "0", font=F_LINK)
    put(ws, tr + 1, 4, f"=B{tr+1}", "0")
    pr = tr + 3
    ws.cell(row=pr, column=1, value="Booth pace before and after the big three sold out").font = F_BOLD
    header(ws, pr + 1, ["Period", "From", "To", "Clicker units", "Units per hour", "Sales"])
    cut = f"{DR}!I{design_row['Avocado']}"
    put(ws, pr + 2, 1, "Creeper, Dumpling or Avocado still on the table")
    put(ws, pr + 2, 2, f"={DR}!B7", TIME_FMT, font=F_LINK)
    put(ws, pr + 2, 3, f"=MAX({DR}!I{design_row['Creeper']},{DR}!I{design_row['Dumpling']},{cut})", TIME_FMT)
    put(ws, pr + 3, 1, "All three gone")
    put(ws, pr + 3, 2, f"=C{pr+2}", TIME_FMT)
    put(ws, pr + 3, 3, f"={DR}!B8", TIME_FMT, font=F_LINK)
    for rr, op1, op2 in ((pr + 2, ">=", "<="), (pr + 3, ">", "<=")):
        put(ws, rr, 4, f'=SUMIFS({rd(C_QTY)},{rd(C_CAT)},"Clickers & Fidget Toys",{rd(C_CLOSE)},"{op1}"&B{rr},{rd(C_CLOSE)},"{op2}"&C{rr})', "0")
        put(ws, rr, 5, f"=D{rr}/((C{rr}-B{rr})*24)", "0.0")
        put(ws, rr, 6, f'=COUNTIFS(Baskets!$C$2:$C${len(baskets)+1},"{op1}"&B{rr},Baskets!$C$2:$C${len(baskets)+1},"{op2}"&C{rr})', "0")
    put(ws, pr + 4, 1, "Some of the drop is the crowd thinning near the 1:00 PM close, so this shows the size of the slump, not how much of it stock caused.", font=F_SUB)
    widths(ws, [44, 11, 10, 10, 11, 10, 12, 10, 10, 12, 12])
    ws.row_dimensions[h].height = 32

    # ------------------------------------------------------------ Event Totals
    ws = wb.create_sheet("Event Totals")
    title(ws, "Event totals: tracked and estimated", "Blue cells are Sam's estimates from the Oct 17 chat. 1 ticket = $1.")
    ET = "'Event Totals'"
    e = EVENT_INPUTS
    rows_in = [
        ("Tickets counted at the end", e["counted"], "Sam: final count."),
        ("Gross tickets (counted + helper food tickets)", e["gross"], "Sam: helpers got about 5 tickets each for food, so close to 900."),
        ("Clickers + copters, low (incl. blowout)", e["cc_low"], "Sam: roughly 500, probably a bit more."),
        ("Clickers + copters, high", e["cc_high"], "Upper end of 'a bit more'. Tracked alone is 482; about 36 leftover clickers went in the blowout."),
        ("Balloon Battle, low", e["battle_low"], "Sam: roughly 50 to 60."),
        ("Balloon Battle, high", e["battle_high"], None),
        ("Plinko tickets per drop", e["plinko_tix_per_drop"], "Sam: 1 ticket per drop."),
        ("Balloon Battle tickets per match (per player)", e["battle_tix_per_match"], "Tracker price."),
        ("CCA's cut", e["cca_cut"], "Tracker README: dollars after CCA's 10%."),
        ("Dollars per ticket", e["dollars_per_ticket"], "Tracker: cash sales are $1 per ticket."),
    ]
    for i, (lab, val, note) in enumerate(rows_in, 4):
        ws.cell(row=i, column=1, value=lab).font = F_BOLD
        inp(ws, i, 2, val, "0%" if lab == "CCA's cut" else "0", note)
    # B4 counted, B5 gross, B6 cc_low, B7 cc_high, B8 b_low, B9 b_high, B10 per drop, B11 per match, B12 cut, B13 $/tix
    h = 15
    header(ws, h, ["Booth", "Tracked tickets", "Estimate low", "Estimate high", "Best estimate",
                   "Goal", "% of goal", "How we got it"])
    clk = f'SUMIFS({rd(C_TIX)},{rd(C_CAT)},"Clickers & Fidget Toys")'
    cop = f'SUMIFS({rd(C_TIX)},{rd(C_CAT)},"Copters")'
    booth_rows = [
        ("Clickers", f"={clk}", f"=B6-B{h+2}", f"=B7-B{h+2}", GOALS["Clickers"],
         "Sam's clickers + copters range minus tracked copters. Includes the untracked blowout."),
        ("Copters", f"={cop}", f"=B{h+2}", f"=B{h+2}", GOALS["Copters"], "All tracked; none in the blowout."),
        ("Balloon Battle", 0, "=B8", "=B9", GOALS["Balloon Battle"], "Sam's estimate; the Battle laptop wasn't exported."),
        ("Plinko", 0, "=B5-B7-B9", "=B5-B6-B8", GOALS["Plinko"], "Inferred: gross minus every other booth. Bucket wasn't counted."),
    ]
    for i, (booth, tracked, lo, hi, goal, how) in enumerate(booth_rows):
        r = h + 1 + i
        put(ws, r, 1, booth, font=F_BOLD)
        put(ws, r, 2, tracked, "0")
        put(ws, r, 3, lo, "0")
        put(ws, r, 4, hi, "0")
        put(ws, r, 5, f"=AVERAGE(C{r}:D{r})", "0", fill=FILL_TOTAL)
        inp(ws, r, 6, goal, "0", "Tracker default goals (status doc).")
        put(ws, r, 7, f"=E{r}/F{r}", "0%")
        put(ws, r, 8, how)
    tr = h + 5
    put(ws, tr, 1, "Total", font=F_BOLD, fill=FILL_TOTAL)
    for col in range(2, 7):
        L = get_column_letter(col)
        put(ws, tr, col, f"=SUM({L}{h+1}:{L}{tr-1})", "0", font=F_BOLD, fill=FILL_TOTAL)
    put(ws, tr, 7, f"=E{tr}/F{tr}", "0%", font=F_BOLD, fill=FILL_TOTAL)
    put(ws, tr, 8, "Best-estimate total equals the gross by construction (Plinko is the remainder).", fill=FILL_TOTAL)
    r0 = tr + 2
    extras = [
        ("Helper food tickets (gross minus counted)", "=B5-B4", "0"),
        ("Gross after CCA's 10% ($)", "=B5*(1-B12)*B13", "$#,##0"),
        ("Plinko drops (best estimate)", f"=E{h+4}/B10", "0"),
        ("Balloon Battle player-matches, low", f"=C{h+3}/B11", "0"),
        ("Balloon Battle player-matches, high", f"=D{h+3}/B11", "0"),
        ("Free Pokemon cards given with purchases", f"=SUM({rd(C_BONUS)})", "0"),
        ("Tracked sales", f"=COUNTA(Baskets!$A$2:$A${len(baskets)+1})", "0"),
        ("Average tickets per tracked sale", f"=SUM(Baskets!$G$2:$G${len(baskets)+1})/B{r0+6}", "0.0"),
    ]
    for i, (lab, f, fmt) in enumerate(extras):
        ws.cell(row=r0 + i, column=1, value=lab).font = F_BOLD
        put(ws, r0 + i, 2, f, fmt)
    widths(ws, [46, 12, 11, 11, 12, 9, 9, 80])
    ws.row_dimensions[h].height = 30

    # ------------------------------------------------------------------ Themes
    ws = wb.create_sheet("Themes")
    title(ws, "Clicker themes", "Theme for each design is set on Design Ranking (column B). Sell-through only counts designs with a known starting count.")
    h = 4
    header(ws, h, ["Theme", "Designs", "Units sold", "Share of clicker units", "Tickets rung up",
                   "Stock (known counts)", "Units (known counts)", "Sell-through"])
    dr_first, dr_last = first, first + len(order) - 1
    DRB = f"{DR}!$B${dr_first}:$B${dr_last}"
    DRC = f"{DR}!$C${dr_first}:$C${dr_last}"
    DRE = f"{DR}!$E${dr_first}:$E${dr_last}"
    DRA = f"{DR}!$A${dr_first}:$A${dr_last}"
    for i, t in enumerate(CLICKER_THEMES):
        r = h + 1 + i
        put(ws, r, 1, t, font=F_BOLD)
        put(ws, r, 2, ", ".join(d for d in order if designs[d]["theme"] == t))
        put(ws, r, 3, f"=SUMIFS({rd(C_QTY)},{rd(C_THEME)},A{r})", "0")
        put(ws, r, 4, f"=C{r}/SUM($C${h+1}:$C${h+4})", "0%")
        put(ws, r, 5, f"=SUMIFS({rd(C_TIX)},{rd(C_THEME)},A{r})", "0")
        put(ws, r, 6, f'=SUMIFS({DRC},{DRB},A{r},{DRA},"<>Oreo")', "0")
        put(ws, r, 7, f'=SUMIFS({DRE},{DRB},A{r},{DRC},"<>",{DRA},"<>Oreo")', "0")
        put(ws, r, 8, f"=IF(F{r}=0,\"\",G{r}/F{r})", "0%", fill=FILL_TOTAL)
    put(ws, h + 5, 1, "Oreo is left out of stock and sell-through (not put out). Egg tart and Burger have no known count, so they count in units but not sell-through.", font=F_SUB)
    # Theme mix in baskets
    r0 = h + 8
    ws.cell(row=r0, column=1, value="How themes mixed in one sale").font = F_BOLD
    header(ws, r0 + 1, ["Sale type", "Sales"])
    put(ws, r0 + 2, 1, "One design")
    put(ws, r0 + 2, 2, f'=COUNTIF(Baskets!$J$2:$J${NB},1)', "0")
    put(ws, r0 + 3, 1, "Several designs, one theme")
    put(ws, r0 + 3, 2, f'=COUNTIFS(Baskets!$J$2:$J${NB},">1",Baskets!$K$2:$K${NB},1)', "0")
    put(ws, r0 + 4, 1, "Several designs, mixed themes")
    put(ws, r0 + 4, 2, f'=COUNTIFS(Baskets!$J$2:$J${NB},">1",Baskets!$K$2:$K${NB},">1")', "0")
    # Theme pair matrix
    r1 = r0 + 7
    ws.cell(row=r1, column=1, value="Sales containing both themes (diagonal: sales with that theme)").font = F_BOLD
    header(ws, r1 + 1, [""] + THEME_ORDER)
    theme_col = {t: get_column_letter(12 + i) for i, t in enumerate(THEME_ORDER)}  # Baskets L..P
    for i, t in enumerate(THEME_ORDER):
        r = r1 + 2 + i
        put(ws, r, 1, t, font=F_BOLD)
        for j, t2 in enumerate(THEME_ORDER):
            a, b = theme_col[t], theme_col[t2]
            put(ws, r, 2 + j, f"=SUMPRODUCT((Baskets!${a}$2:${a}${NB}>0)*(Baskets!${b}$2:${b}${NB}>0))", "0",
                fill=FILL_BAND if i == j else None)
    widths(ws, [30, 52, 10, 12, 12, 12, 12, 12])
    ws.row_dimensions[h].height = 30

    # ---------------------------------------------------------------- Pairings
    ws = wb.create_sheet("Pairings")
    title(ws, "Designs bought together", "Each cell counts sales that had both designs. The diagonal counts sales with that design at all.")
    pair_designs = [d for d in order if d != "Oreo"] + ["Full kit", "Kit w/ toroidal"]
    design_col = {}
    for i, d in enumerate([d for d in STARTING_STOCK] + ["Full kit", "Kit w/ toroidal"]):
        design_col[d] = get_column_letter(17 + i)  # Baskets Q onward
    h = 4
    ws.cell(row=h, column=1, value="Top pairs").font = F_BOLD
    header(ws, h + 1, ["Design A", "Design B", "Sales together"])
    top = [p for p in m["pairs"] if p[2] >= 2]
    for i, (a, b, _) in enumerate(top):
        r = h + 2 + i
        put(ws, r, 1, a)
        put(ws, r, 2, b)
        ca, cb = design_col[a], design_col[b]
        put(ws, r, 3, f"=SUMPRODUCT((Baskets!${ca}$2:${ca}${NB}>0)*(Baskets!${cb}$2:${cb}${NB}>0))", "0")
    r1 = h + 3 + len(top)
    put(ws, r1, 1, "Pairs that happened once are in the matrix below.", font=F_SUB)
    r2 = r1 + 2
    ws.cell(row=r2, column=1, value="Pair matrix").font = F_BOLD
    header(ws, r2 + 1, [""] + pair_designs)
    for i, d in enumerate(pair_designs):
        r = r2 + 2 + i
        put(ws, r, 1, d, font=F_BOLD)
        for j, d2 in enumerate(pair_designs):
            a, b = design_col[d], design_col[d2]
            put(ws, r, 2 + j, f"=SUMPRODUCT((Baskets!${a}$2:${a}${NB}>0)*(Baskets!${b}$2:${b}${NB}>0))", '0;-0;""',
                fill=FILL_BAND if i == j else None)
    widths(ws, [22, 22] + [9] * len(pair_designs))
    ws.row_dimensions[r2 + 1].height = 45
    for j in range(len(pair_designs)):
        ws.cell(row=r2 + 1, column=2 + j).alignment = Alignment(wrap_text=True, vertical="center")

    # ------------------------------------------------------------ Demographics
    ws = wb.create_sheet("Demographics")
    title(ws, "Who bought what (units, by who the item was for)",
          "Tags are optional on the tracker. Untagged = total minus tagged. Small groups (Adults: 4 sales) are hints, not findings.")
    segs = [("Kid", C_AGE), ("Teen", C_AGE), ("Adult", C_AGE), ("Male", C_GEN), ("Female", C_GEN)]
    h = 4
    header(ws, h, ["Design", "Total units", "Kid", "Teen", "Adult", "Untagged age",
                   "Male", "Female", "Untagged gender", "Boys' share (tagged)"])
    for i, d in enumerate(pair_designs):
        r = h + 1 + i
        put(ws, r, 1, d, font=F_BOLD)
        put(ws, r, 2, f"=SUMIFS({rd(C_QTY)},{rd(C_ITEM)},A{r})", "0")
        cols = {}
        for j, (seg, col) in enumerate(segs[:3]):
            put(ws, r, 3 + j, f'=SUMIFS({rd(C_QTY)},{rd(C_ITEM)},$A{r},{rd(col)},"{seg}")', '0;-0;""')
        put(ws, r, 6, f"=B{r}-SUM(C{r}:E{r})", '0;-0;""')
        for j, (seg, col) in enumerate(segs[3:]):
            put(ws, r, 7 + j, f'=SUMIFS({rd(C_QTY)},{rd(C_ITEM)},$A{r},{rd(col)},"{seg}")', '0;-0;""')
        put(ws, r, 9, f"=B{r}-G{r}-H{r}", '0;-0;""')
        put(ws, r, 10, f'=IF(G{r}+H{r}=0,"",G{r}/(G{r}+H{r}))', "0%")
    tr = h + 1 + len(pair_designs)
    put(ws, tr, 1, "Total", font=F_BOLD, fill=FILL_TOTAL)
    for col in range(2, 10):
        L = get_column_letter(col)
        put(ws, tr, col, f"=SUM({L}{h+1}:{L}{tr-1})", "0", font=F_BOLD, fill=FILL_TOTAL)
    put(ws, tr, 10, f'=G{tr}/(G{tr}+H{tr})', "0%", font=F_BOLD, fill=FILL_TOTAL)
    # Theme share within each segment
    r0 = tr + 3
    ws.cell(row=r0, column=1, value="Theme share of each group's units").font = F_BOLD
    header(ws, r0 + 1, ["Theme", "Kid", "Teen", "Adult", "Male", "Female", "All"])
    for i, t in enumerate(THEME_ORDER):
        r = r0 + 2 + i
        put(ws, r, 1, t, font=F_BOLD)
        for j, (seg, col) in enumerate(segs):
            L = get_column_letter(2 + j)
            put(ws, r, 2 + j, f'=SUMIFS({rd(C_QTY)},{rd(C_THEME)},$A{r},{rd(col)},"{seg}")/SUMIFS({rd(C_QTY)},{rd(col)},"{seg}")', "0%")
        put(ws, r, 7, f"=SUMIFS({rd(C_QTY)},{rd(C_THEME)},$A{r})/SUM({rd(C_QTY)})", "0%")
    r1 = r0 + 2 + len(THEME_ORDER) + 1
    ws.cell(row=r1, column=1, value="Sales by tag").font = F_BOLD
    header(ws, r1 + 1, ["Tag", "Sales", "Share of sales"])
    tags = [("Kid", "E"), ("Teen", "E"), ("Adult", "E"), ("Male", "F"), ("Female", "F")]
    for i, (tag, col) in enumerate(tags):
        r = r1 + 2 + i
        put(ws, r, 1, tag, font=F_BOLD)
        put(ws, r, 2, f'=COUNTIF(Baskets!${col}$2:${col}${NB},"{tag}")', "0")
        put(ws, r, 3, f"=B{r}/COUNTA(Baskets!$A$2:$A${NB})", "0%")
    put(ws, r1 + 7, 1, "Age untagged")
    put(ws, r1 + 7, 2, f"=COUNTA(Baskets!$A$2:$A${NB})-SUM(B{r1+2}:B{r1+4})", "0")
    put(ws, r1 + 7, 3, f"=B{r1+7}/COUNTA(Baskets!$A$2:$A${NB})", "0%")
    put(ws, r1 + 8, 1, "Gender untagged")
    put(ws, r1 + 8, 2, f"=COUNTA(Baskets!$A$2:$A${NB})-SUM(B{r1+5}:B{r1+6})", "0")
    put(ws, r1 + 8, 3, f"=B{r1+8}/COUNTA(Baskets!$A$2:$A${NB})", "0%")
    widths(ws, [22, 10, 8, 8, 8, 10, 8, 8, 10, 12])
    ws.row_dimensions[h].height = 30

    # ---------------------------------------------------------------- Timeline
    ws = wb.create_sheet("Timeline")
    title(ws, "Sales by half hour", "By the time each sale was closed. Official hours were 10:30 AM to 1:00 PM; the table sold from 10:01 AM.")
    h = 4
    header(ws, h, ["From", "To", "Sales", "Tickets", "Units"])
    t0 = datetime(2026, 10, 17, 10, 0)
    for i in range(8):
        r = h + 1 + i
        put(ws, r, 1, t0 + timedelta(minutes=30 * i), TIME_FMT)
        put(ws, r, 2, f"=A{r}+TIME(0,30,0)", TIME_FMT)
        crit = f'Baskets!$C$2:$C${NB},">="&A{r},Baskets!$C$2:$C${NB},"<"&B{r}'
        put(ws, r, 3, f"=COUNTIFS({crit})", "0")
        put(ws, r, 4, f"=SUMIFS(Baskets!$G$2:$G${NB},{crit})", "0")
        put(ws, r, 5, f"=SUMIFS(Baskets!$H$2:$H${NB},{crit})", "0")
    tr = h + 9
    put(ws, tr, 1, "Total", font=F_BOLD, fill=FILL_TOTAL)
    put(ws, tr, 2, "", fill=FILL_TOTAL)
    for col in (3, 4, 5):
        L = get_column_letter(col)
        put(ws, tr, col, f"=SUM({L}{h+1}:{L}{tr-1})", "0", font=F_BOLD, fill=FILL_TOTAL)
    r0 = tr + 2
    ws.cell(row=r0, column=1, value="Checkout speed (sale opened to Done)").font = F_BOLD
    put(ws, r0 + 1, 1, "Median seconds")
    put(ws, r0 + 1, 2, f"=MEDIAN(Baskets!$D$2:$D${NB})", "0")
    put(ws, r0 + 2, 1, "Fastest seconds")
    put(ws, r0 + 2, 2, f"=MIN(Baskets!$D$2:$D${NB})", "0")
    put(ws, r0 + 3, 1, "Ticket totals per sale")
    header(ws, r0 + 4, ["Tickets", "Sales"])
    for i, t in enumerate(sorted(m["baskets"]["tickets_hist"])):
        r = r0 + 5 + i
        put(ws, r, 1, t, "0")
        put(ws, r, 2, f"=COUNTIF(Baskets!$G$2:$G${NB},A{r})", "0")
    widths(ws, [24, 12, 9, 9, 9])

    # ----------------------------------------------------------------- Copters
    ws = wb.create_sheet("Copters")
    title(ws, "Copters: this year vs last year", "Same prices both years. Last year's number is net of CCA's 10% and food tickets (Sam).")
    ws["A4"] = "Last year, net tickets"
    ws["A4"].font = F_BOLD
    inp(ws, 4, 2, EVENT_INPUTS["copters_last_year_net"], "0", "Sam: roughly 450 after the 10% and after tickets pulled to feed people.")
    ws["A5"] = "This year, tracked tickets"
    ws["A5"].font = F_BOLD
    put(ws, 5, 2, f'=SUMIFS({rd(C_TIX)},{rd(C_CAT)},"Copters")', "0")
    ws["A6"] = "This year, net of CCA's 10%"
    ws["A6"].font = F_BOLD
    put(ws, 6, 2, f"=B5*(1-{ET}!B12)", "0")
    ws["A7"] = "Change"
    ws["A7"].font = F_BOLD
    put(ws, 7, 2, "=B6/B4-1", "0%", font=F_BOLD, fill=FILL_TOTAL)
    ws["A8"] = "Copter share of the sales table's tracked tickets"
    ws["A8"].font = F_BOLD
    put(ws, 8, 2, f"=B5/SUM({rd(C_TIX)})", "0%")
    header(ws, 10, ["Kit", "Kits sold", "Tickets"])
    for i, k in enumerate(["Full kit", "Kit w/ toroidal"]):
        r = 11 + i
        put(ws, r, 1, k, font=F_BOLD)
        put(ws, r, 2, f"=SUMIFS({rd(C_QTY)},{rd(C_ITEM)},A{r})", "0")
        put(ws, r, 3, f"=SUMIFS({rd(C_TIX)},{rd(C_ITEM)},A{r})", "0")
    header(ws, 14, ["Part", "Starting stock", "Used", "Left for next year"])
    uses = {"Launchers": "=B11+B12", "Propellers": "=B11", "Toroidal propellers": "=B12", "Ripcords": "=B11+B12"}
    for i, (p, s) in enumerate(COPTER_PARTS.items()):
        r = 15 + i
        put(ws, r, 1, p, font=F_BOLD)
        inp(ws, r, 2, s, "0", "Tracker built-in count (Oct 16)." if i == 0 else None)
        put(ws, r, 3, uses[p], "0")
        put(ws, r, 4, f"=B{r}-C{r}", "0", fill=FILL_TOTAL)
    widths(ws, [46, 14, 10, 18])

    # ----------------------------------------------------------------- Baskets
    ws = wb.create_sheet("Baskets")
    all_items = [d for d in STARTING_STOCK] + ["Full kit", "Kit w/ toroidal"]
    hdr = (["Record", "Opened", "Closed", "Checkout seconds", "Age", "Gender", "Tickets",
            "Units", "Free cards", "Designs in sale", "Themes in sale"] + THEME_ORDER + all_items)
    header(ws, 1, hdr)
    for i, b in enumerate(sorted(baskets, key=lambda b: int(b["record"][1:]))):
        r = 2 + i
        put(ws, r, 1, b["record"])
        put(ws, r, 2, b["opened"], DT_FMT)
        put(ws, r, 3, b["closed"], DT_FMT)
        put(ws, r, 4, f"=(C{r}-B{r})*86400", "0")
        put(ws, r, 5, b["age"] or None)
        put(ws, r, 6, b["gender"] or None)
        put(ws, r, 7, f"=SUMIFS({rd(C_TIX)},{rd(C_REC)},A{r})", "0")
        put(ws, r, 8, f"=SUMIFS({rd(C_QTY)},{rd(C_REC)},A{r})", "0")
        put(ws, r, 9, f"=SUMIFS({rd(C_BONUS)},{rd(C_REC)},A{r})", "0")
        dfirst, dlast = get_column_letter(17), get_column_letter(16 + len(all_items))
        put(ws, r, 10, f'=COUNTIF({dfirst}{r}:{dlast}{r},">0")', "0")
        put(ws, r, 11, f'=COUNTIF(L{r}:P{r},">0")', "0")
        for j, t in enumerate(THEME_ORDER):
            put(ws, r, 12 + j, f'=SUMIFS({rd(C_QTY)},{rd(C_REC)},$A{r},{rd(C_THEME)},"{t}")', '0;-0;""')
        for j, d in enumerate(all_items):
            put(ws, r, 17 + j, f'=SUMIFS({rd(C_QTY)},{rd(C_REC)},$A{r},{rd(C_ITEM)},"{d}")', '0;-0;""')
    ws.freeze_panes = "B2"
    widths(ws, [8, 18, 18, 9, 7, 8, 8, 7, 7, 8, 8] + [10] * (len(THEME_ORDER) + len(all_items)))
    ws.row_dimensions[1].height = 45

    # ---------------------------------------------------------------- Raw Data
    ws = wb.create_sheet("Raw Data")
    header(ws, 1, header_cols + ["theme", "qty_for_rate"])
    for i, r in enumerate(rows):
        rr = 2 + i
        for j, hname in enumerate(header_cols):
            v = r[hname]
            if hname in ("opened_at", "closed_at", "exported_at"):
                v = parse_dt(v)
                c = ws.cell(row=rr, column=j + 1, value=v)
                c.number_format = DT_FMT
            elif hname in ("qty", "item_tickets", "record_total_tickets"):
                c = ws.cell(row=rr, column=j + 1, value=int(v))
            elif hname == "bonus_cards":
                c = ws.cell(row=rr, column=j + 1, value=int(v) if v != "" else None)
            else:
                c = ws.cell(row=rr, column=j + 1, value=v if v != "" else None)
            c.font = F_BODY
        c = ws.cell(row=rr, column=len(header_cols) + 1,
                    value=f'=IFERROR(INDEX({DRB},MATCH({C_ITEM}{rr},{DRA},0)),IF({C_CAT}{rr}="Copters","Copters","Other"))')
        c.font = F_BODY
        c = ws.cell(row=rr, column=len(header_cols) + 2, value=f"=MIN({C_QTY}{rr},{DR}!$B$6)")
        c.font = F_BODY
    ws.freeze_panes = "A2"
    for j in range(len(header_cols) + 2):
        ws.column_dimensions[get_column_letter(j + 1)].width = 14
    ws.column_dimensions[C_ITEM].width = 22
    ws.column_dimensions[C_CAT].width = 22

    fill_leaderboard(wsL, m, design_row, rd, C_QTY, C_CAT, C_AGE, C_GEN, NB)
    # Findings text (static; written from the metrics)
    fill_findings(wsF, m)
    return wb


def fill_leaderboard(ws, m, design_row, rd, C_QTY, C_CAT, C_AGE, C_GEN, NB):
    DR = "'Design Ranking'"
    title(ws, "Leaderboard", "Ranked by sales rate: clickers sold per hour while the design was on the table. Fair across uneven stock.")
    h = 4
    header(ws, h, ["Rate rank", "Design", "Sales rate (units/hr)", "Sold out at", "Sell-through",
                   "Units sold", "Starting stock", "Blended rank (70% rate)", "Notes"])
    for i, d in enumerate(m["ranking_order"]):
        r = h + 1 + i
        s = design_row[d]
        put(ws, r, 1, f"={DR}!R{s}", "0", font=F_BOLD, fill=FILL_TOTAL)
        put(ws, r, 2, f"={DR}!A{s}", font=F_BOLD)
        put(ws, r, 3, f"={DR}!L{s}", "0.0", font=F_LINK)
        put(ws, r, 4, f'=IF({DR}!H{s}="Unknown","Unknown",IF({DR}!I{s}="","Didn\'t",{DR}!I{s}))', TIME_FMT, font=F_LINK)
        put(ws, r, 5, f'=IF({DR}!G{s}="","Unknown",{DR}!G{s})', "0%", font=F_LINK)
        put(ws, r, 6, f"={DR}!E{s}", "0", font=F_LINK)
        put(ws, r, 7, f'=IF({DR}!C{s}="","Restocked",{DR}!C{s})', "0", font=F_LINK)
        put(ws, r, 8, f"={DR}!P{s}", "0", font=F_LINK)
        put(ws, r, 9, f'=IF({DR}!S{s}="","",{DR}!S{s})', font=F_LINK)
    r0 = h + 2 + len(m["ranking_order"])
    put(ws, r0, 1, "Not ranked: Oreo (probably never put out) and Burger (not in the stock count, 1 sold).", font=F_SUB)

    r1 = r0 + 2
    ws.cell(row=r1, column=1, value="Most popular with each group").font = Font(name=FONT, size=11, bold=True, color=RED)
    header(ws, r1 + 1, ["Group", "", "Sales", "Clicker units", "Top designs (number of that group's sales that picked it)"])
    for i, g in enumerate(m["top_by_group"]):
        r = r1 + 2 + i
        put(ws, r, 1, g["group"], font=F_BOLD)
        put(ws, r, 2, "")
        crit_b, crit_r = [], []
        if g["age"]:
            crit_b.append(f'Baskets!$E$2:$E${NB},"{g["age"]}"')
            crit_r.append(f'{rd(C_AGE)},"{g["age"]}"')
        if g["gender"]:
            crit_b.append(f'Baskets!$F$2:$F${NB},"{g["gender"]}"')
            crit_r.append(f'{rd(C_GEN)},"{g["gender"]}"')
        put(ws, r, 3, f"=COUNTIFS({','.join(crit_b)})", "0")
        put(ws, r, 4, f'=SUMIFS({rd(C_QTY)},{rd(C_CAT)},"Clickers & Fidget Toys",{",".join(crit_r)})', "0")
        put(ws, r, 5, g["top"], note=("Top-designs text is written by the build script from Raw Data; the counts to its left are live formulas."
                                       if i == 0 else None))
    rr = r1 + 2 + len(m["top_by_group"])
    put(ws, rr + 1, 1, "Untagged sales aren't in any group. Groups under about 5 sales are hints, not findings. Next year, split Kid into Elementary, Middle, and High school.", font=F_SUB)
    widths(ws, [10, 22, 12, 12, 11, 10, 11, 12, 60])
    ws.column_dimensions["E"].width = 11
    ws.row_dimensions[h].height = 32
    # Group table text lives in column E; let it spill over F:I visually
    for i in range(len(m["top_by_group"])):
        ws.merge_cells(start_row=r1 + 2 + i, start_column=5, end_row=r1 + 2 + i, end_column=9)
    ws.merge_cells(start_row=r1 + 1, start_column=5, end_row=r1 + 1, end_column=9)
    ws.merge_cells(start_row=r1 + 1, start_column=1, end_row=r1 + 1, end_column=2)
    for i in range(len(m["top_by_group"])):
        ws.merge_cells(start_row=r1 + 2 + i, start_column=1, end_row=r1 + 2 + i, end_column=2)


def fill_findings(ws, m):
    title(ws, "Findings", "Conclusions are backed by the tracked data. Hypotheses are guesses worth testing next year.")
    r = 4
    for head, items in findings_sections(m):
        c = ws.cell(row=r, column=1, value=head)
        c.font = Font(name=FONT, size=11, bold=True, color=RED)
        r += 1
        for it in items:
            c = ws.cell(row=r, column=1, value="-")
            c.font = F_BODY
            c.alignment = Alignment(vertical="top")
            c = ws.cell(row=r, column=2, value=it)
            c.font = F_BODY
            c.alignment = WRAP
            r += 1
        r += 1
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 130


def findings_sections(m):
    d = m["designs"]
    ev = m["event"]
    lost = m["lost"]
    pace = m["pace"]
    sections = [
        ("Headline numbers", [
            f"About {ev['gross']:.0f} tickets gross (830 counted + about {ev['helper_food']:.0f} given to helpers for food), about ${ev['net_after_cca']:.0f} after CCA's 10%. Goal was {ev['goal_total']:,}: we hit about {ev['gross']/ev['goal_total']:.0%}.",
            f"Tracked: {m['tracked']['sales']} sales, {m['tracked']['clicker_units']} clickers and {m['tracked']['copter_kits']} copter kits, {m['tracked']['tickets']} tickets.",
            f"Plinko (inferred): about {ev['plinko_low']:.0f} to {ev['plinko_high']:.0f} tickets, so roughly {ev['plinko_drops_mid']:.0f} drops at 1 ticket each. That's about {ev['plinko_mid']/GOALS['Plinko']:.1f}x its 100-ticket goal and the second-biggest booth.",
        ]),
        ("Best designs (by sales rate: units per hour on the table)", [
            "Rate rank: " + ", ".join(f"{d[x]['rate_rank']}. {x} ({d[x]['rate']:.1f}/hr)" for x in m['ranking_order'][:6]) + ". Egg tart and Chocolate bar tie at 5.",
            f"Dumpling and Salmon nigiri were the real demand leaders: about 7 an hour each while on the table. They also happened to be the two designs we made the most of, which is why they lead on units too.",
            f"Creeper sold out first. All 10 were gone by 11:48 AM with no bulk buys, about 90 minutes before the last sale. Its rate ({d['Creeper']['rate']:.1f}/hr) says we needed about {d['Creeper']['rate'] * m['window']['hours']:.0f} to last the day ({d['Creeper']['ideal_next_year']:.0f} with a 20% buffer), not 10.",
            "Spam musubi only made 6 and sold its last 2 at 1:20 PM. It sold out, but slowly (1.8/hr): it was right-sized, not a hidden hit.",
            "Slowest: S'mores (0.9/hr) and LEGO brick (0.6/hr).",
            ("The blended rank (70% rate, 30% units) comes out in the same order, so the call doesn't depend on how you weigh it."
             if m["blended_order"] == m["ranking_order"] else
             "The blended rank (70% rate, 30% units) shuffles a few spots; see Design Ranking."),
        ]),
        ("What we ran short on", [
            f"Creeper, Dumpling, and Avocado sold out before the end. At their own pace, unlimited stock would have sold about {lost['units']:.0f} more clickers, about {lost['tickets']:.0f} tickets at {m['avg_tix_per_clicker']:.2f} tix each. Creeper alone is {d['Creeper']['lost_units']:.0f} of those.",
            f"That's a floor. Clicker pace fell from {pace['before_units_per_hour']:.0f} units/hour to {pace['after_units_per_hour']:.0f} after {pace['cutoff'][:5]} PM, once all three were gone. The crowd was thinning toward the 1:00 PM close too, so we can't pin the whole drop on stock.",
            f"Meanwhile {lost['left_over_units']:.0f} clickers were left for the blowout: Hot dog bun 12, WD-40 7, S'mores 6, Chocolate bar 4, Oreo 4, LEGO 2, Creeper cube 1. We printed about the right number of clickers in the wrong mix.",
        ]),
        ("Themes and pairings", [
            f"Every Asian food design with a known count (Salmon nigiri, Dumpling, Spam musubi) sold out; Egg tart was restocked, so its count is unknown. Minecraft sold {m['themes']['Minecraft']['sell_through_known']:.0%} of stock. Snacks & treats sold {m['themes']['Snacks & treats']['sell_through_known']:.0%} (Avocado carried it). Workshop sold {m['themes']['Workshop']['sell_through_known']:.0%}.",
            f"Variety beat sets. Of the {m['baskets']['multi_unit']} sales with 2 or more items, {m['baskets']['multi_unit'] - m['baskets']['same_design_multi']} mixed designs and only {m['baskets']['same_design_multi']} were several of one design. {m['baskets']['multi_theme']} sales mixed themes.",
            "The most common theme mix was Asian food + Snacks & treats (12 sales).",
            "Top pairs: Avocado + Dumpling (4 sales); Creeper + Dumpling, Avocado + Egg tart, Dumpling + Egg tart (3 each).",
            f"The mixers: {m['mixers'][0][0]} showed up in {m['mixers'][0][1]} mixed sales alongside {m['mixers'][0][2]} different designs, and {m['mixers'][1][0]} in {m['mixers'][1][1]}. They're the anchors kids built a 3-for-10 around.",
        ]),
        ("Who bought what (counted by customers, so one bulk buyer doesn't skew it)", [
            "Kids were the core: 32 of the 45 age-tagged sales. Teens 9, Adults 4.",
        ] + [f"{g['group']}: {g['top']}." for g in m['top_by_group'] if g['group'] in
             ("Kids", "Teens", "Kid boys", "Kid girls", "Boys and men", "Girls and women")] + [
            "Creeper and WD-40 skewed boys (7 of 8 gender-tagged units each). Egg tart skewed girls (7 of 8).",
            "Teens went for Salmon nigiri: 6 of 9 teen sales included one.",
            "Every gender-tagged copter buyer was a boy.",
            "Kid lumps elementary and middle schoolers into one bucket, which hides a lot. Next year: Elementary, Middle, and High school.",
        ]),
        ("Copters", [
            f"{m['tracked']['copter_kits']} kits for {m['tracked']['copter_tickets']} tickets, against roughly 450 net last year at the same prices. That's down about {abs(ev['copter_change']):.0%}.",
            f"2 of the 5 kits were toroidal. {m['copters']['parts_left']['Launchers']} launchers, {m['copters']['parts_left']['Propellers']} propellers, {m['copters']['parts_left']['Toroidal propellers']} toroidal props, and {m['copters']['parts_left']['Ripcords']} ripcords carry over. Don't buy more parts.",
        ]),
        ("Hypotheses to test", [
            "Clickers pulled the same kids and tickets copters used to get: they took 94% of the sales table's tickets. Test: sell copters at their own table or with a flying demo.",
            "Families from last year already own a copter. Test: ask at the table, or sell replacement props and ripcords to returning kids.",
            "Egg tart may not have reached the table until about 11:00 AM (first sale 11:07, about 45 minutes after most designs started selling). If so, it sold 11 in about two hours, ranks 4th on rate, and deserves more stock.",
            "LEGO brick sold twice in two minutes at 10:41 AM, then never again. It may have been buried behind other designs.",
            "The 3-for-10 deal plus a free Pokémon card seems to pull kids up to 3: more sales stopped at 3 clickers (20) than at 2 (12).",
        ]),
        ("Next year", [
            "Rebalance the print queue with the Ideal stock column on Stock & Lost Sales (this year's rate x 3.3 hours + 20%): " + ", ".join(
                f"{x} {d[x]['ideal_next_year']:.0f} (had {d[x]['stock']})" for x in ("Dumpling", "Salmon nigiri", "Creeper", "Avocado", "Hot dog bun", "S'mores")) + ".",
            "Get every booth onto the tracker. Plinko and Battle together were about 43% of the money, and we have no data on them.",
            "Count stock again the morning of, after any restock, and ring the blowout up as its own item so it shows up.",
            "Swap the Kid / Teen / Adult tags for Elementary / Middle / High school / Adult, and tag every sale (about a third were untagged).",
        ]),
    ]
    return sections


def write_summary_csv(m, path):
    """One-tab values snapshot for Google Drive (the connector only takes text)."""
    d = m["designs"]
    ev = m["event"]
    out = []
    out.append(["Ram-a-thon 2026 results: values snapshot (v1.0, 2026-10-17)"])
    out.append(["Live version with formulas (13 tabs): ramathon_2026_results_v1.0.xlsx on the ramathon-carnival branch, tools/ramathon_sales_tracker/results_2026/. Drag it into Drive to get it as a Google Sheet."])
    out.append([])
    out.append(["LEADERBOARD: ranked by sales rate (units per hour while on the table)"])
    out.append(["Rate rank", "Design", "Theme", "Sales rate (units/hr)", "Units sold", "Starting stock",
                "Sell-through", "Sold out at", "Lost units (sold out early)", "Lost tickets",
                "Ideal stock next year", "Left over", "Blended rank (70% rate)"])
    for x in m["ranking_order"]:
        y = d[x]
        out.append([y["rate_rank"], x, y["theme"], round(y["rate"], 1), y["units"],
                    y["stock"] if y["stock"] is not None else "Restocked (unknown)",
                    f"{y['sell_through']:.0%}" if y["sell_through"] is not None else "Unknown",
                    (datetime.strptime(y["sold_out_at"], "%H:%M").strftime("%-I:%M %p") if y["sold_out_at"]
                     else ("Unknown" if y["stock"] is None else "Didn't")),
                    round(y["lost_units"], 1), round(y["lost_tickets"]),
                    int(y["ideal_next_year"]), "" if y["left_over"] is None else y["left_over"],
                    y["blended_rank"]])
    out.append(["", "Total lost", "", "", "", "", "", "", round(m["lost"]["units"], 1), round(m["lost"]["tickets"])])
    out.append(["Not ranked: Oreo (probably never put out, 4 left), Burger (not in the stock count, 1 sold)."])
    out.append([])
    out.append(["EVENT TOTALS (1 ticket = $1)"])
    out.append(["Booth", "Tracked tickets", "Estimate low", "Estimate high", "Best estimate", "Goal", "% of goal", "How"])
    cl, ch = EVENT_INPUTS["cc_low"] - ev["copters"], EVENT_INPUTS["cc_high"] - ev["copters"]
    rows_ev = [
        ("Clickers", m["tracked"]["clicker_tickets"], cl, ch, ev["clickers_mid"], GOALS["Clickers"], "Sam's 500 to 530 for clickers + copters, minus copters. Includes the blowout."),
        ("Copters", ev["copters"], ev["copters"], ev["copters"], ev["copters"], GOALS["Copters"], "All tracked."),
        ("Balloon Battle", 0, EVENT_INPUTS["battle_low"], EVENT_INPUTS["battle_high"], ev["battle_mid"], GOALS["Balloon Battle"], "Sam's estimate (about 10 to 12 player-matches)."),
        ("Plinko", 0, ev["plinko_low"], ev["plinko_high"], ev["plinko_mid"], GOALS["Plinko"], "Inferred: 900 gross minus every other booth. About 330 drops at 1 tix."),
    ]
    for b, t, lo, hi, mid, g, how in rows_ev:
        out.append([b, t, lo, hi, round(mid), g, f"{mid/g:.0%}", how])
    out.append(["Total", m["tracked"]["tickets"], "", "", ev["gross"], ev["goal_total"], f"{ev['gross']/ev['goal_total']:.0%}", f"830 counted + about {ev['helper_food']} helper food tickets. About ${ev['net_after_cca']:.0f} after CCA's 10%."])
    out.append([])
    out.append(["MOST POPULAR WITH EACH GROUP (number of that group's sales that picked the design)"])
    out.append(["Group", "Top designs"])
    for g in m["top_by_group"]:
        out.append([g["group"], g["top"]])
    out.append([])
    out.append(["THEMES"])
    out.append(["Theme", "Units sold", "Sell-through (known counts)"])
    for t in CLICKER_THEMES:
        th = m["themes"][t]
        out.append([t, th["units"], f"{th['sell_through_known']:.0%}"])
    out.append([])
    out.append(["TOP PAIRS (sales with both)"])
    out.append(["Design A", "Design B", "Sales"])
    for a, b, n in m["pairs"]:
        if n >= 3:
            out.append([a, b, n])
    out.append([])
    for head, items in findings_sections(m):
        out.append([])
        out.append([head.upper()])
        for it in items:
            out.append(["", it])
    with open(path, "w", newline="") as f:
        csv.writer(f).writerows(out)


def jsonable(o):
    if isinstance(o, datetime):
        return o.isoformat()
    raise TypeError(type(o))


def main():
    src, out_xlsx, out_json = sys.argv[1:4]
    header_cols, rows = load(src)
    # The export's battle, prize, and constant columns are empty for the sales
    # laptop; the Raw Data tab keeps the columns with data (full CSV is in the repo).
    header_cols = [h for h in header_cols if h in RAW_KEEP]
    m, baskets = compute(rows)
    wb = build(header_cols, rows, m, baskets, out_xlsx)
    wb.save(out_xlsx)
    with open(out_json, "w") as f:
        json.dump(m, f, indent=2, default=jsonable)
    if len(sys.argv) > 4:
        write_summary_csv(m, sys.argv[4])
    print("ok", out_xlsx, out_json)


if __name__ == "__main__":
    main()

# CHANGELOG
# v1.0 (2026-10-17): Initial release. Ranking, lost sales, themes, pairings, demographics, timeline, copters, event totals with Plinko inferred.
